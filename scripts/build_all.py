#!/usr/bin/env python3
"""Build the workbook HTML, DOCX, print PDF, and reader EPUB from Markdown.

Canonical sources are ``frontmatter/*.md``, ``chapters/**/*.md``, and
``signature/**/*.md``. ``python scripts/build_all.py --proof`` also includes
``build/fixtures/typesetting-proof.md``, which is not a chapter.

``journal/private/`` is never read. Pandoc and Typst are not used.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import re
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
HTML_PATH = ROOT / "build" / "workbook.html"
DOCX_PATH = ROOT / "build" / "The_Closer_Project_Illustrated_Workbook_v1_draft.docx"
PDF_PATH = ROOT / "build" / "The_Closer_Project_Illustrated_Workbook_v1_draft.pdf"
EPUB_PATH = ROOT / "build" / "The_Closer_Project_Reader_v1_draft.epub"
FIXTURE_PATH = ROOT / "build" / "fixtures" / "typesetting-proof.md"
PDF_ERROR_PATH = ROOT / "build" / "pdf-print-error.txt"
WORKBOOK_CSS = ROOT / "styles" / "workbook.css"
EPUB_CSS = ROOT / "styles" / "epub.css"
PRIVATE_ROOT = (ROOT / "journal" / "private").resolve()
STAMP = "2026-10-14T00:00:00Z"
CONTENT_TWIPS = 8122  # 5.64 in of text on a 7 in page

REQUIRED_HEADINGS = [
    "The hook",
    "Why it might matter",
    "Before you begin",
    "The experience",
    "Three ways to play",
    "Going deeper",
    "Variations and replay",
    "Examples",
    "Afterward",
    "Our Story",
    "Notes",
]

INDEX_SPECS = (
    ("index-alphabetical", "Alphabetical index"),
    ("index-time", "Time"),
    ("index-emotional-intensity", "Emotional intensity"),
    ("index-energy", "Energy"),
    ("index-format", "Format"),
    ("index-mood", "Mood"),
    ("index-revisit", "Revisit"),
)

TIME_ORDER = ["5-10", "10-20", "20-45", "45-90", "90-180", "multi-day"]
INTENSITY_ORDER = ["light", "personal", "deep", "intense"]
ENERGY_ORDER = ["low", "medium", "high"]

JOURNAL_LABELS = {
    "session": "Session",
    "partner_a": "Private page · one of you",
    "partner_b": "Private page · the other",
    "shared": "Shared page",
    "revisit": "Revisit",
}

BROWSERS = [
    Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
    Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
    Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
]

INK = RGBColor(0x24, 0x1F, 0x1C)
COPPER = RGBColor(0xA8, 0x5A, 0x32)
BURGUNDY = RGBColor(0x7C, 0x3A, 0x46)
TEAL = RGBColor(0x3E, 0x64, 0x64)
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
HYPERLINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"

INLINE_RE = re.compile(
    r"`([^`]+)`"
    r"|!\[([^\]]*)\]\(([^)\s]+)\)"
    r"|\[([^\]]+)\]\(([^)\s]+)\)"
    r"|\*\*([^*]+)\*\*"
    r"|\*([^*]+)\*"
)
H2_RE = re.compile(r"^##\s+(.+?)\s*$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*$")
FENCE_OPEN_RE = re.compile(r"^```([A-Za-z0-9_-]*)\s*$")
FENCE_CLOSE_RE = re.compile(r"^```\s*$")
HR_RE = re.compile(r"^(\*\s*\*\s*\*|-{3,}|_{3,})\s*$")
LIST_RE = re.compile(r"^(\s*)(\d+\.|[-*+])\s+(.*)$")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")


@dataclass
class SourceDoc:
    path: Path
    rel: str
    kind: str
    meta: dict
    body: str
    blocks: list
    anchor: str
    title: str
    heading_ids: list[str | None] = field(default_factory=list)


@dataclass
class Piece:
    kind: str
    anchor: str
    title: str
    inner: str
    toc_group: str | None = None
    group_anchor: str | None = None


@dataclass
class IndexSection:
    anchor: str
    title: str
    blocks: list


@dataclass
class ImageSink:
    notes: list[str] = field(default_factory=list)
    found: dict[str, Path] = field(default_factory=dict)

    def place(self, src: str, alt: str, missing_note: bool) -> str:
        rel = src.strip().replace("\\", "/")
        path = resolve_asset(rel)
        if path is None:
            note = f"SKIP image, file not on disk: {rel}"
            if note not in self.notes:
                self.notes.append(note)
            if missing_note:
                return f'<p class="motif-note">No image placed. The file is not on disk: {esc(rel)}</p>'
            return ""
        posix = path.resolve().relative_to(ROOT.resolve()).as_posix()
        self.found[posix] = path
        return f'<img src="asset:{esc(posix)}" alt="{esc(alt)}"/>'


class Bookmarks:
    def __init__(self) -> None:
        self.by_html: dict[str, str] = {}
        self.used: set[str] = set()
        self.next_id = 1

    def register(self, html_id: str) -> str:
        if html_id in self.by_html:
            return self.by_html[html_id]
        base = re.sub(r"[^A-Za-z0-9_.-]", "", html_id) or "section"
        if not base[0].isalpha():
            base = "a" + base
        name = base[:40]
        n = 2
        while name in self.used:
            suffix = str(n)
            name = base[: 40 - len(suffix)] + suffix
            n += 1
        self.used.add(name)
        self.by_html[html_id] = name
        return name

    def add(self, paragraph, html_id: str) -> None:
        name = self.by_html[html_id]
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self.next_id))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self.next_id))
        self.next_id += 1
        p_el = paragraph._p
        p_pr = p_el.find(qn("w:pPr"))
        if p_pr is not None:
            p_pr.addnext(start)
        else:
            p_el.insert(0, start)
        p_el.append(end)


def configure_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, OSError):
            pass


def esc(text: str) -> str:
    return html.escape(str(text), quote=True)


def slugify(text: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return slug or "section"


def is_private_path(path: Path) -> bool:
    parts = [part.lower() for part in path.parts]
    for index in range(len(parts) - 1):
        if parts[index] == "journal" and parts[index + 1] == "private":
            return True
    try:
        resolved = path.resolve()
    except OSError:
        return True
    return resolved == PRIVATE_ROOT or PRIVATE_ROOT in resolved.parents


def assert_readable(path: Path) -> None:
    if is_private_path(path):
        raise SystemExit(f"Refusing to read private journal path: {path}")


def within_root(path: Path) -> Path:
    resolved = path.resolve()
    root = ROOT.resolve()
    if resolved != root and root not in resolved.parents:
        raise SystemExit(f"Refusing path outside the project: {path}")
    assert_readable(resolved)
    return resolved


def assert_sources_unchanged() -> None:
    ledger = ROOT / "source" / "SOURCE_CHECKSUMS.txt"
    if not ledger.is_file():
        raise SystemExit(f"Missing checksum ledger: {ledger}")
    for line in ledger.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        digest, name = line.split()
        path = ROOT / "source" / name
        if not path.is_file():
            raise SystemExit(f"Missing source original: {name}")
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        if actual != digest:
            raise SystemExit(
                f"Source original changed and will not be overwritten or rebuilt from: {name}"
            )


def heading_problems(body: str) -> list[str]:
    found = [match.group(1).strip() for match in (H2_RE.match(line) for line in body.splitlines()) if match]
    missing: list[str] = []
    position = 0
    for required in REQUIRED_HEADINGS:
        try:
            position = found.index(required, position) + 1
        except ValueError:
            missing.append(required)
    return missing


def rel_posix(path: Path) -> str:
    return path.resolve().relative_to(ROOT.resolve()).as_posix()


def collect_paths(proof: bool) -> list[Path]:
    def markdown_in(folder: str, pattern: str) -> list[Path]:
        base = ROOT / folder
        if not base.is_dir():
            return []
        found = [path for path in base.glob(pattern) if path.is_file() and path.name.lower() != "readme.md"]
        return sorted((within_root(path) for path in found), key=rel_posix)

    paths = markdown_in("frontmatter", "*.md")
    paths.extend(markdown_in("chapters", "**/*.md"))
    paths.extend(markdown_in("signature", "**/*.md"))
    if proof:
        if not FIXTURE_PATH.is_file():
            raise SystemExit(f"Missing typesetting fixture: {FIXTURE_PATH}")
        paths.append(within_root(FIXTURE_PATH))
    if not paths:
        raise SystemExit("No markdown sources found. No files were written.")
    return paths


def split_frontmatter(text: str, origin: str) -> tuple[dict, str]:
    if not text.startswith("---"):
        return {}, text
    lines = text.splitlines()
    if lines[0].strip() != "---":
        return {}, text
    for index in range(1, len(lines)):
        if lines[index].strip() == "---":
            raw = "\n".join(lines[1:index])
            body = "\n".join(lines[index + 1 :]).lstrip("\n")
            try:
                loaded = yaml.safe_load(raw) if raw.strip() else {}
            except yaml.YAMLError as exc:
                raise SystemExit(f"{origin}: YAML frontmatter could not be parsed: {exc}") from exc
            if loaded is None:
                loaded = {}
            if not isinstance(loaded, dict):
                raise SystemExit(f"{origin}: YAML frontmatter must be a mapping. No files were written.")
            return loaded, body
    raise SystemExit(f"{origin}: YAML frontmatter is missing its closing ---. No files were written.")


def kind_of(rel: str) -> str:
    if rel == FIXTURE_PATH.resolve().relative_to(ROOT.resolve()).as_posix():
        return "fixture"
    top = rel.split("/", 1)[0]
    if top == "frontmatter":
        return "front"
    if top == "chapters":
        return "chapter"
    if top == "signature":
        return "signature"
    raise SystemExit(f"Unexpected markdown source: {rel}")


def first_h1(body: str) -> str | None:
    for line in body.splitlines():
        match = re.match(r"^#\s+(.+?)\s*$", line)
        if match:
            return match.group(1).strip()
    return None


def chapter_anchor(meta: dict, origin: str) -> str:
    if "id" not in meta:
        raise SystemExit(f"{origin} is missing id in frontmatter. No files were written.")
    raw = str(meta["id"]).strip()
    if not raw.isdigit():
        raise SystemExit(f"{origin} has a non-numeric id {raw!r}. No files were written.")
    return f"chapter-{int(raw):03d}"


def make_anchor(kind: str, meta: dict, rel: str, origin: str) -> str:
    if kind == "chapter":
        return chapter_anchor(meta, origin)
    if kind == "fixture":
        return "fixture-not-a-chapter"
    if kind == "signature":
        # Ids such as S01 are not chapter numbers. Do not call chapter_anchor.
        token = str(meta.get("id") or Path(rel).stem)
        return "signature-" + slugify(token)[:32]
    return "front-" + slugify(Path(rel).stem)


def make_title(meta: dict, body: str, path: Path) -> str:
    if meta.get("title"):
        return str(meta["title"]).strip()
    found = first_h1(body)
    return found or path.stem


def is_table_sep(line: str) -> bool:
    return TABLE_SEP_RE.match(line) is not None


def split_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def starts_block(line: str) -> bool:
    if HEADING_RE.match(line) or FENCE_OPEN_RE.match(line) or HR_RE.match(line) or LIST_RE.match(line):
        return True
    return line.lstrip().startswith("|")


def parse_blocks(body: str, origin: str) -> list:
    lines = body.splitlines()
    index = 0
    blocks = []
    total = len(lines)
    while index < total:
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        fence = FENCE_OPEN_RE.match(line)
        if fence:
            lang = fence.group(1).lower()
            index += 1
            buf: list[str] = []
            while index < total and not FENCE_CLOSE_RE.match(lines[index]):
                buf.append(lines[index])
                index += 1
            if index >= total:
                raise SystemExit(f"{origin}: unclosed {lang or 'code'} fence. No files were written.")
            index += 1
            blocks.append(("fence", lang, "\n".join(buf)))
            continue
        heading = HEADING_RE.match(line)
        if heading:
            blocks.append(("heading", len(heading.group(1)), heading.group(2).strip()))
            index += 1
            continue
        if HR_RE.match(line):
            blocks.append(("hr",))
            index += 1
            continue
        if line.lstrip().startswith("|") and index + 1 < total and is_table_sep(lines[index + 1]):
            header = split_row(line)
            index += 2
            rows = []
            while index < total and lines[index].strip().startswith("|"):
                rows.append(split_row(lines[index]))
                index += 1
            blocks.append(("table", header, rows))
            continue
        listed = LIST_RE.match(line)
        if listed:
            ordered = listed.group(2)[0].isdigit()
            items = []
            while index < total:
                item = LIST_RE.match(lines[index])
                if not item:
                    break
                this_ordered = item.group(2)[0].isdigit()
                if this_ordered != ordered:
                    break
                items.append(item.group(3).strip())
                index += 1
            blocks.append(("list", ordered, items))
            continue
        buf = [line.strip()]
        index += 1
        while index < total and lines[index].strip() and not starts_block(lines[index]):
            buf.append(lines[index].strip())
            index += 1
        blocks.append(("para", " ".join(buf)))
    return blocks


def parse_journal(text: str) -> list[dict]:
    sections: list[dict] = []
    current: dict | None = None
    for raw in text.splitlines():
        if not raw.strip():
            continue
        if raw.startswith("- "):
            if current is None:
                current = {"key": "notes", "intro": "", "items": []}
                sections.append(current)
            current["items"].append(raw[2:].strip())
            continue
        match = re.match(r"^([A-Za-z_][A-Za-z0-9_]*)\s*:\s*(.*)$", raw)
        if match:
            current = {"key": match.group(1), "intro": match.group(2).strip(), "items": []}
            sections.append(current)
            continue
        if current is None:
            current = {"key": "notes", "intro": raw.strip(), "items": []}
            sections.append(current)
        else:
            current["intro"] = (current["intro"] + " " + raw.strip()).strip()
    return sections


def parse_cards(text: str) -> list[tuple[str, str]]:
    cards = []
    for raw in text.splitlines():
        if not raw.strip():
            continue
        if "|" in raw:
            title, body = raw.split("|", 1)
            cards.append((title.strip(), body.strip()))
        else:
            cards.append(("", raw.strip()))
    return cards


def body_blocks(doc: SourceDoc) -> list:
    blocks = doc.blocks
    if doc.kind in {"chapter", "signature", "fixture"} and blocks:
        first = blocks[0]
        if first[0] == "heading" and first[1] == 1 and first[2] == doc.title:
            return blocks[1:]
    return blocks


def load_doc(path: Path) -> SourceDoc:
    assert_readable(path)
    rel = rel_posix(path)
    text = path.read_text(encoding="utf-8-sig")
    meta, body = split_frontmatter(text, rel)
    kind = kind_of(rel)
    if kind == "chapter" and "title" not in meta:
        raise SystemExit(f"{rel} is missing title in frontmatter. No files were written.")
    return SourceDoc(
        path=path,
        rel=rel,
        kind=kind,
        meta=meta,
        body=body,
        blocks=parse_blocks(body, rel),
        anchor=make_anchor(kind, meta, rel, rel),
        title=make_title(meta, body, path),
    )


def prepare(docs: list[SourceDoc]) -> dict[str, str]:
    """Assign unique ids. Returns part number -> part opener anchor."""
    used: set[str] = set()
    part_anchors: dict[str, str] = {}
    for doc in docs:
        if doc.anchor in used:
            raise SystemExit(f"Duplicate anchor {doc.anchor} from {doc.rel}. No files were written.")
        used.add(doc.anchor)
        if doc.kind != "chapter":
            continue
        part = doc.meta.get("part")
        if part is None:
            continue
        key = str(part)
        if key in part_anchors:
            continue
        anchor = f"part-{int(key):02d}" if key.isdigit() else "part-" + slugify(key)
        if anchor in used:
            raise SystemExit(f"Duplicate anchor {anchor}. No files were written.")
        used.add(anchor)
        part_anchors[key] = anchor
    for doc in docs:
        visible = body_blocks(doc)
        mapping: list[str | None] = []
        for block in visible:
            if block[0] != "heading":
                mapping.append(None)
                continue
            base = f"{doc.anchor}-{slugify(block[2])}"
            heading_id = base
            suffix = 2
            while heading_id in used:
                heading_id = f"{base}-{suffix}"
                suffix += 1
            used.add(heading_id)
            mapping.append(heading_id)
        doc.heading_ids = mapping
    return part_anchors


def check_chapters(docs: list[SourceDoc]) -> None:
    problems = []
    for doc in docs:
        if doc.kind != "chapter":
            continue
        missing = heading_problems(doc.body)
        if missing:
            problems.append(
                f"{doc.rel} is missing required headings or has them out of order: " + "; ".join(missing)
            )
    if problems:
        joined = "\n".join(problems)
        raise SystemExit(joined + "\nNo workbook files were written.")


def check_fixture(docs: list[SourceDoc]) -> None:
    fixtures = [doc for doc in docs if doc.kind == "fixture"]
    if len(fixtures) != 1:
        raise SystemExit("The proof build needs exactly one typesetting fixture. No files were written.")
    fixture = fixtures[0]
    if fixture.title != "FIXTURE NOT A CHAPTER":
        raise SystemExit("The typesetting fixture must be titled FIXTURE NOT A CHAPTER. No files were written.")
    languages = [block[1] for block in fixture.blocks if block[0] == "fence"]
    kinds = [block[0] for block in fixture.blocks]
    if "journal" not in languages or "cards" not in languages or "table" not in kinds:
        raise SystemExit("The fixture must include a journal block, a card block, and a table. No files were written.")
    if not fixture.meta.get("motif"):
        raise SystemExit("The fixture needs a motif path so a missing image can be skipped. No files were written.")


def signature_problems(doc: SourceDoc) -> list[str]:
    """Same heading order as a chapter, plus a motif path. S01 is not a chapter id."""
    if doc.kind != "signature":
        return []
    problems = []
    missing = heading_problems(doc.body)
    if missing:
        problems.append(
            f"{doc.rel} is missing required headings or has them out of order: " + "; ".join(missing)
        )
    motif = doc.meta.get("motif")
    if not isinstance(motif, str) or not motif.strip():
        problems.append(f"{doc.rel} is missing a motif path.")
    return problems


def check_signatures(docs: list[SourceDoc]) -> None:
    problems = []
    for doc in docs:
        problems.extend(signature_problems(doc))
    if problems:
        raise SystemExit("\n".join(problems) + "\nNo workbook files were written.")


def load_bibliography() -> SourceDoc | None:
    path = ROOT / "research" / "BIBLIOGRAPHY.md"
    if not path.is_file():
        return None
    assert_readable(path)
    rel = rel_posix(path)
    text = path.read_text(encoding="utf-8-sig")
    meta, body = split_frontmatter(text, rel)
    return SourceDoc(
        path=path,
        rel=rel,
        kind="bibliography",
        meta=meta,
        body=body,
        blocks=parse_blocks(body, rel),
        anchor="bibliography",
        title=make_title(meta, body, path),
    )


def load_sources(proof: bool) -> tuple[list[SourceDoc], dict[str, str]]:
    docs = [load_doc(path) for path in collect_paths(proof)]
    bibliography = load_bibliography()
    if bibliography is not None:
        docs.append(bibliography)
    part_anchors = prepare(docs)
    check_chapters(docs)
    check_signatures(docs)
    if proof:
        check_fixture(docs)
    return docs, part_anchors


def parse_inline(text: str) -> list[tuple]:
    runs = []
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            runs.append(("text", text[position : match.start()]))
        if match.group(1) is not None:
            runs.append(("code", match.group(1)))
        elif match.group(2) is not None:
            runs.append(("image", match.group(2), match.group(3)))
        elif match.group(4) is not None:
            runs.append(("link", match.group(4), match.group(5)))
        elif match.group(6) is not None:
            runs.append(("bold", match.group(6)))
        else:
            runs.append(("italic", match.group(7) or ""))
        position = match.end()
    if position < len(text):
        runs.append(("text", text[position:]))
    return runs


def normalize_href(href: str) -> str:
    href = href.strip()
    if href.lower().startswith("javascript:"):
        return "#"
    if re.fullmatch(r"0*\d{1,3}", href):
        return f"#chapter-{int(href):03d}"
    return href


def inline_xhtml(text: str, images: ImageSink | None) -> str:
    parts = []
    for run in parse_inline(text):
        kind = run[0]
        if kind == "text":
            parts.append(esc(run[1]))
        elif kind == "bold":
            parts.append(f"<strong>{esc(run[1])}</strong>")
        elif kind == "italic":
            parts.append(f"<em>{esc(run[1])}</em>")
        elif kind == "code":
            parts.append(f"<code>{esc(run[1])}</code>")
        elif kind == "link":
            href = normalize_href(run[2])
            parts.append(f'<a href="{esc(href)}">{esc(run[1])}</a>')
        elif kind == "image":
            if images is None:
                parts.append(esc(run[1] or run[2]))
            else:
                parts.append(images.place(run[2], run[1] or "Illustration", missing_note=False))
    return "".join(parts)


def writing_lines(count: int) -> str:
    return '<div class="writing-lines" aria-hidden="true">' + ("<span></span>" * count) + "</div>"


def cover_heading(title: str) -> bool:
    low = title.strip().lower()
    return (
        low.startswith("cover this")
        or low.startswith("sample answer")
        or low.startswith("how rowan and vale closed")
        or low.startswith("facilitator section")
    )


def closes_player_text(title: str) -> bool:
    """The heading after a facilitator key starts a new sheet, so the key does not share a page with play."""
    return title.strip().lower().startswith("facilitator section")


def journal_label(key: str) -> str:
    return JOURNAL_LABELS.get(key, key.replace("_", " ").capitalize())


def journal_xhtml(text: str) -> str:
    parts = [
        '<section class="journal-spread">',
        '<p class="journal-kicker">Our Story</p>',
        "<p class=\"journal-private\">Sharing either private page is optional.</p>",
    ]
    for section in parse_journal(text):
        parts.append('<div class="prompt-block">')
        parts.append(f"<h3>{esc(journal_label(section['key']))}</h3>")
        if section["intro"]:
            parts.append(f'<p class="journal-instruction">{inline_xhtml(section["intro"], None)}</p>')
        if section["items"]:
            for item in section["items"]:
                parts.append(f'<p class="prompt">{inline_xhtml(item, None)}</p>')
                parts.append(writing_lines(8))
        else:
            parts.append(writing_lines(2 if section["key"] == "session" else 8))
        parts.append("</div>")
    parts.append("</section>")
    return "\n".join(parts)


def cards_xhtml(text: str) -> str:
    parts = ['<div class="card-grid">']
    for title, body in parse_cards(text):
        parts.append('<article class="card">')
        if title:
            parts.append(f"<h3>{inline_xhtml(title, None)}</h3>")
        if body:
            parts.append(f"<p>{inline_xhtml(body, None)}</p>")
        parts.append("</article>")
    parts.append("</div>")
    return "\n".join(parts)


def table_xhtml(header: list[str], rows: list[list[str]]) -> str:
    parts = ["<table><thead><tr>"]
    for cell in header:
        parts.append(f"<th>{inline_xhtml(cell, None)}</th>")
    parts.append("</tr></thead><tbody>")
    width = len(header)
    for row in rows:
        padded = list(row) + [""] * max(0, width - len(row))
        parts.append("<tr>")
        for cell in padded[:width]:
            parts.append(f"<td>{inline_xhtml(cell, None)}</td>")
        parts.append("</tr>")
    parts.append("</tbody></table>")
    return "".join(parts)


def list_xhtml(ordered: bool, items: list[str]) -> str:
    tag = "ol" if ordered else "ul"
    body = "".join(f"<li>{inline_xhtml(item, None)}</li>" for item in items)
    return f"<{tag}>{body}</{tag}>"


def meta_line(doc: SourceDoc) -> str:
    meta = doc.meta
    bits = []
    if meta.get("classification"):
        bits.append(str(meta["classification"]).replace("-", " "))
    if meta.get("primary_minutes"):
        bits.append(f"{meta['primary_minutes']} minutes")
    elif meta.get("time_band"):
        bits.append(f"{meta['time_band']} minutes")
    if meta.get("emotional_intensity"):
        bits.append(str(meta["emotional_intensity"]))
    if meta.get("energy"):
        bits.append(f"{meta['energy']} energy")
    if meta.get("location"):
        bits.append(str(meta["location"]))
    return " · ".join(bits)


def resolve_asset(src: str) -> Path | None:
    if not src or "://" in src or src.startswith("mailto:"):
        return None
    candidate = within_root(ROOT / src)
    if candidate.is_file():
        return candidate
    return None


def opener_xhtml(doc: SourceDoc, images: ImageSink, anchors: set[str]) -> str:
    bits = ['<header class="chapter-open">']
    if doc.kind == "fixture":
        bits.append('<p class="kicker">Fixture · not a catalog entry</p>')
        bits.append('<p class="chapter-number">Proof</p>')
    elif doc.kind == "signature":
        bits.append('<p class="kicker">Signature session</p>')
        if doc.meta.get("id"):
            bits.append(f'<p class="chapter-number">{esc(doc.meta["id"])}</p>')
    else:
        part = doc.meta.get("part")
        part_title = str(doc.meta.get("part_title") or "")
        if part and part_title:
            bits.append(f'<p class="kicker">Part {esc(part)} · {esc(part_title)}</p>')
        elif part_title:
            bits.append(f'<p class="kicker">{esc(part_title)}</p>')
        elif part:
            bits.append(f'<p class="kicker">Part {esc(part)}</p>')
        if doc.meta.get("id") is not None:
            raw = str(doc.meta["id"])
            shown = raw.zfill(3) if raw.isdigit() else raw
            bits.append(f'<p class="chapter-number">{esc(shown)}</p>')
    bits.append(f"<h1>{esc(doc.title)}</h1>")
    if doc.meta.get("deck"):
        bits.append(f'<p class="deck">{esc(doc.meta["deck"])}</p>')
    motif = doc.meta.get("motif")
    if motif:
        placed = images.place(str(motif), f"Motif for {doc.title}", missing_note=doc.kind == "fixture")
        if placed.startswith("<p"):
            bits.append(placed)
        elif placed:
            bits.append(f'<figure class="motif">{placed}<figcaption>Motif</figcaption></figure>')
    summary = meta_line(doc)
    if summary:
        bits.append(f'<p class="meta-line">{esc(summary)}</p>')
    links = doc.meta.get("cross_links") or []
    if isinstance(links, list) and links:
        rendered = []
        for item in links:
            raw = str(item).strip()
            if raw.isdigit():
                anchor = f"chapter-{int(raw):03d}"
                label = raw.zfill(3)
                if anchor in anchors:
                    rendered.append(f'<a href="#{esc(anchor)}">{esc(label)}</a>')
                else:
                    rendered.append(esc(label))
            else:
                rendered.append(esc(raw))
        bits.append('<p class="see-also">See also ' + ", ".join(rendered) + "</p>")
    bits.append("</header>")
    return "\n".join(bits)


def blocks_xhtml(doc: SourceDoc, images: ImageSink, anchors: set[str]) -> str:
    parts = []
    visible = body_blocks(doc)
    close_next = False
    for block, heading_id in zip(visible, doc.heading_ids):
        kind = block[0]
        if kind == "heading":
            level = min(max(block[1], 1), 4)
            cover = " page-break" if cover_heading(block[2]) or close_next else ""
            close_next = closes_player_text(block[2])
            parts.append(f'<h{level} id="{esc(heading_id)}" class="{cover.strip()}">{esc(block[2])}</h{level}>' if cover else f'<h{level} id="{esc(heading_id)}">{esc(block[2])}</h{level}>')
        elif kind == "para":
            image_only = parse_inline(block[1])
            if len(image_only) == 1 and image_only[0][0] == "image":
                alt, src = image_only[0][1], image_only[0][2]
                placed = images.place(src, alt or "Illustration", missing_note=True)
                if placed:
                    parts.append(placed if placed.startswith("<p") else f"<p>{placed}</p>")
            else:
                parts.append(f"<p>{inline_xhtml(block[1], images)}</p>")
        elif kind == "list":
            parts.append(list_xhtml(block[1], block[2]))
        elif kind == "table":
            parts.append(table_xhtml(block[1], block[2]))
        elif kind == "fence":
            lang, content = block[1], block[2]
            if lang == "journal":
                parts.append(journal_xhtml(content))
            elif lang == "cards":
                parts.append(cards_xhtml(content))
            else:
                parts.append(f"<pre><code>{esc(content)}</code></pre>")
        elif kind == "hr":
            parts.append("<hr/>")
    return "\n".join(parts)


def article_xhtml(doc: SourceDoc, images: ImageSink, anchors: set[str]) -> str:
    css = {
        "front": "front-piece",
        "chapter": "chapter",
        "signature": "chapter signature",
        "fixture": "chapter fixture",
        "bibliography": "endmatter",
    }[doc.kind]
    parts = [f'<article class="{css}" id="{esc(doc.anchor)}">']
    if doc.kind == "fixture":
        parts.append(
            '<p class="proof-banner">FIXTURE NOT A CHAPTER. This is a typesetting sample, not a catalog entry.</p>'
        )
    if doc.kind in {"chapter", "signature", "fixture"}:
        parts.append(opener_xhtml(doc, images, anchors))
    parts.append(blocks_xhtml(doc, images, anchors))
    parts.append("</article>")
    return "\n".join(part for part in parts if part)


def title_inner(proof: bool, images: ImageSink) -> str:
    banner = ""
    if proof:
        banner = (
            '<p class="proof-banner">Typesetting proof. The fixture at the end is not a chapter '
            "and is not catalog entry 001.</p>"
        )
    cover = images.place(
        "assets/covers/cover-front.svg",
        "Front cover of The Closer Project on a midnight ground.",
        missing_note=False,
    )
    figure = f'<figure class="cover">{cover}</figure>' if cover and not cover.startswith("<p") else ""
    return f'''<header class="title-page" id="title-page">
<p class="kicker">Draft v1 · 14 October 2026</p>
{figure}
<h1>The Closer Project</h1>
<p class="subtitle">Illustrated workbook</p>
<div class="rule"></div>
<p class="deck">Experiences for two adults who already like each other.</p>
{banner}
</header>'''


def toc_inner(docs: list[SourceDoc], part_anchors: dict[str, str]) -> str:
    items = [
        '<nav class="toc" id="contents">',
        '<p class="kicker">Draft v1</p>',
        "<h1>Contents</h1>",
    ]
    current = None
    ol_open = False

    def close_ol() -> None:
        nonlocal ol_open
        if ol_open:
            items.append("</ol>")
            ol_open = False

    def begin(label: str, href: str | None) -> None:
        nonlocal current, ol_open
        close_ol()
        current = label
        if href:
            items.append(f'<p class="toc-group"><a href="{esc(href)}">{esc(label)}</a></p>')
        else:
            items.append(f'<p class="toc-group">{esc(label)}</p>')
        items.append("<ol>")
        ol_open = True

    for doc in docs:
        if doc.kind == "front":
            label, href = "Front matter", None
        elif doc.kind == "chapter":
            part = doc.meta.get("part")
            part_title = str(doc.meta.get("part_title") or "")
            label = f"Part {part}" + (f" · {part_title}" if part_title else "") if part else "Chapters"
            key = str(part) if part is not None else ""
            href = f"#{part_anchors[key]}" if key in part_anchors else None
        elif doc.kind == "signature":
            label, href = "Signature sessions", None
        elif doc.kind == "bibliography":
            continue
        else:
            label, href = "Typesetting fixture (not a chapter)", None
        if label != current:
            begin(label, href)
        items.append(f'<li><a href="#{esc(doc.anchor)}">{esc(doc.title)}</a></li>')
    close_ol()
    items.append('<p class="toc-group">Indexes</p>')
    items.append("<ol>")
    for anchor, title in INDEX_SPECS:
        items.append(f'<li><a href="#{esc(anchor)}">{esc(title)}</a></li>')
    items.append("</ol>")
    for doc in docs:
        if doc.kind == "bibliography":
            items.append(f'<p class="toc-group"><a href="#{esc(doc.anchor)}">{esc(doc.title)}</a></p>')
    items.append("</nav>")
    return "\n".join(items)


def colophon_inner(docs: list[SourceDoc], notes: list[str], images: ImageSink) -> str:
    sources = "\n".join(f"<li>{esc(doc.rel)}</li>" for doc in docs)
    if notes:
        skipped = "<ul>\n" + "\n".join(f"<li>{esc(note)}</li>" for note in notes) + "\n</ul>"
    else:
        skipped = "<p>No referenced image was missing.</p>"
    cover = images.place(
        "assets/covers/cover-back.svg",
        "Back cover of The Closer Project.",
        missing_note=False,
    )
    figure = f'<figure class="cover">{cover}</figure>' if cover and not cover.startswith("<p") else ""
    return f'''<section class="colophon" id="colophon">
{figure}
<h1>Build note</h1>
<p>Set 14 October 2026 from Markdown. This file is a draft. The folder <code>journal/private/</code> is not read.</p>
<p>Sources in this file:</p>
<ul>
{sources}
</ul>
{skipped}
</section>'''


def part_opener_path(part: object) -> str | None:
    text = str(part).strip()
    if not text.isdigit():
        return None
    return f"assets/parts/{int(text):02d}-opener.svg"


def part_introductions() -> dict[str, str]:
    path = ROOT / "parts" / "INTRODUCTIONS.md"
    if not path.is_file():
        return {}
    found: dict[str, str] = {}
    current: str | None = None
    buf: list[str] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        match = re.match(r"^##\s+(\d+)\s*$", line.strip())
        if match:
            if current and buf:
                found[current] = " ".join(buf).strip()
            current = str(int(match.group(1)))
            buf = []
        elif current and line.strip() and not line.startswith("#"):
            buf.append(line.strip())
    if current and buf:
        found[current] = " ".join(buf).strip()
    return found


def part_inner(part: str, part_title: str, anchor: str, images: ImageSink) -> str:
    title_text = (part_title or f"Part {part}").strip()
    figure = ""
    src = part_opener_path(part)
    if src:
        placed = images.place(src, title_text, missing_note=False)
        if placed and not placed.startswith("<p"):
            figure = f'<figure class="part-art">{placed}</figure>'
    intro = part_introductions().get(str(int(part)) if str(part).isdigit() else "", "")
    prose = f'<p class="part-intro">{esc(intro)}</p>' if intro else ""
    return (
        f'<section class="part-opener" id="{esc(anchor)}">'
        f'<p class="kicker">Part {esc(part)}</p>'
        f"{figure}"
        f"<h1>{esc(title_text)}</h1>"
        f"{prose}</section>"
    )


def string_list(value) -> list[str]:
    if value is None:
        return []
    items = value if isinstance(value, (list, tuple)) else [value]
    found = []
    for item in items:
        text = str(item).strip()
        if text:
            found.append(text)
    return found


def index_refs(docs: list[SourceDoc]) -> list[dict]:
    """Chapter frontmatter only, using anchors already assigned.

    Does not call chapter_anchor. Signature ids such as S01 stay out of these indexes.
    """
    refs = []
    for doc in docs:
        if doc.kind != "chapter":
            continue
        meta = doc.meta
        raw = str(meta.get("id", "")).strip()
        refs.append(
            {
                "anchor": doc.anchor,
                "title": doc.title,
                "shown": raw.zfill(3) if raw.isdigit() else raw,
                "sort_id": int(raw) if raw.isdigit() else 10**6,
                "time_band": str(meta.get("time_band") or "").strip(),
                "emotional_intensity": str(meta.get("emotional_intensity") or "").strip(),
                "energy": str(meta.get("energy") or "").strip(),
                "formats": string_list(meta.get("formats", meta.get("format"))),
                "moods": string_list(meta.get("moods", meta.get("mood"))),
                "repeatable": meta.get("repeatable") is True,
            }
        )
    return refs


def group_blocks(refs: list[dict], values_fn, order: list[str] | None) -> list:
    buckets: dict[str, list] = {}
    labels: dict[str, str] = {}
    for ref in refs:
        values = [value for value in values_fn(ref) if value]
        if not values:
            values = ["(not set)"]
        for value in values:
            key = value.casefold()
            labels.setdefault(key, value)
            buckets.setdefault(key, []).append(ref)

    def sorter(key: str):
        if order:
            folded = [item.casefold() for item in order]
            if key in folded:
                return (0, folded.index(key), key)
        if key == "(not set)":
            return (2, 0, key)
        return (1, 0, key)

    blocks = []
    for key in sorted(buckets, key=sorter):
        members = sorted(buckets[key], key=lambda ref: (ref["sort_id"], ref["title"].casefold()))
        deduped = []
        seen = set()
        for ref in members:
            if ref["anchor"] in seen:
                continue
            seen.add(ref["anchor"])
            deduped.append(ref)
        blocks.append(("group", labels[key], deduped))
    return blocks


def build_index_sections(docs: list[SourceDoc]) -> list[IndexSection]:
    refs = index_refs(docs)
    alpha = sorted(refs, key=lambda ref: (ref["title"].casefold(), ref["sort_id"], ref["anchor"]))
    revisit = sorted(
        (ref for ref in refs if ref["repeatable"]),
        key=lambda ref: (ref["sort_id"], ref["title"].casefold()),
    )
    revisit_blocks: list = [("rows", revisit)] if revisit else [("note", "No chapter is marked repeatable.")]
    return [
        IndexSection("index-alphabetical", "Alphabetical index", [("rows", alpha)]),
        IndexSection("index-time", "Time", group_blocks(refs, lambda ref: [ref["time_band"]], TIME_ORDER)),
        IndexSection(
            "index-emotional-intensity",
            "Emotional intensity",
            group_blocks(refs, lambda ref: [ref["emotional_intensity"]], INTENSITY_ORDER),
        ),
        IndexSection("index-energy", "Energy", group_blocks(refs, lambda ref: [ref["energy"]], ENERGY_ORDER)),
        IndexSection("index-format", "Format", group_blocks(refs, lambda ref: ref["formats"], None)),
        IndexSection("index-mood", "Mood", group_blocks(refs, lambda ref: ref["moods"], None)),
        IndexSection("index-revisit", "Revisit", revisit_blocks),
    ]


def index_html(section: IndexSection) -> str:
    parts = [
        f'<section class="index" id="{esc(section.anchor)}">',
        '<p class="kicker">Index</p>',
        f"<h1>{esc(section.title)}</h1>",
        '<p class="index-note">Generated from the chapter frontmatter at build time.</p>',
    ]
    for block in section.blocks:
        kind = block[0]
        if kind == "note":
            parts.append(f"<p>{esc(block[1])}</p>")
        elif kind == "rows":
            parts.append('<table class="index-table"><thead><tr><th>Title</th><th>ID</th></tr></thead><tbody>')
            for ref in block[1]:
                parts.append(
                    "<tr>"
                    f'<td><a href="#{esc(ref["anchor"])}">{esc(ref["title"])}</a></td>'
                    f"<td>{esc(ref['shown'])}</td>"
                    "</tr>"
                )
            parts.append("</tbody></table>")
        elif kind == "group":
            parts.append(f"<h2>{esc(block[1])}</h2>")
            parts.append("<ul>")
            for ref in block[2]:
                label = f"{ref['shown']}  {ref['title']}".strip()
                parts.append(f'<li><a href="#{esc(ref["anchor"])}">{esc(label)}</a></li>')
            parts.append("</ul>")
    parts.append("</section>")
    return "\n".join(parts)


def build_index_pieces(docs: list[SourceDoc]) -> list[Piece]:
    pieces = []
    first_anchor = INDEX_SPECS[0][0]
    for section in build_index_sections(docs):
        pieces.append(
            Piece(
                "index",
                section.anchor,
                section.title,
                index_html(section),
                "Indexes",
                first_anchor if section.anchor == first_anchor else None,
            )
        )
    return pieces


def build_pieces(docs: list[SourceDoc], part_anchors: dict[str, str], proof: bool, images: ImageSink) -> list[Piece]:
    anchors = {doc.anchor for doc in docs}
    pieces = [
        Piece("title", "title-page", "The Closer Project", title_inner(proof, images)),
        Piece("toc", "contents", "Contents", toc_inner(docs, part_anchors)),
    ]
    seen_parts: set[str] = set()
    group_for = {
        "front": "Front matter",
        "signature": "Signature sessions",
        "fixture": "Typesetting fixture (not a chapter)",
        "bibliography": "Bibliography",
    }
    bibliography = None
    for doc in docs:
        if doc.kind == "bibliography":
            bibliography = doc
            continue
        group = group_for.get(doc.kind)
        group_anchor = None
        if doc.kind == "chapter":
            part = doc.meta.get("part")
            part_title = str(doc.meta.get("part_title") or "")
            key = str(part) if part is not None else ""
            group = f"Part {part}" + (f" · {part_title}" if part_title else "") if part else "Chapters"
            group_anchor = part_anchors.get(key)
            if key and key not in seen_parts and key in part_anchors:
                seen_parts.add(key)
                pieces.append(
                    Piece(
                        "part",
                        part_anchors[key],
                        group,
                        part_inner(str(part), part_title, part_anchors[key], images),
                        group,
                        part_anchors[key],
                    )
                )
        pieces.append(Piece(doc.kind, doc.anchor, doc.title, article_xhtml(doc, images, anchors), group, group_anchor))
    pieces.extend(build_index_pieces(docs))
    if bibliography is not None:
        pieces.append(
            Piece(
                "bibliography",
                bibliography.anchor,
                bibliography.title,
                article_xhtml(bibliography, images, anchors),
                "Bibliography",
                bibliography.anchor,
            )
        )
    pieces.append(Piece("colophon", "colophon", "Build note", colophon_inner(docs, images.notes, images)))
    return pieces


def html_document(pieces: list[Piece], css: str) -> str:
    body = "\n".join(piece.inner.replace('src="asset:', 'src="../') for piece in pieces)
    title = "The Closer Project — Illustrated Workbook (draft v1)"
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<meta name="generator" content="scripts/build_all.py"/>
<title>{esc(title)}</title>
<!-- styles/workbook.css, inlined so print does not depend on a second file -->
<style>
{css}
</style>
</head>
<body>
{body}
</body>
</html>
"""


def xhtml_document(title: str, body: str, css_href: str) -> str:
    return f"""<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="en" xml:lang="en">
<head>
<title>{esc(title)}</title>
<link rel="stylesheet" type="text/css" href="{esc(css_href)}"/>
</head>
<body>
{body}
</body>
</html>
"""


def epub_nav(pieces: list[Piece]) -> str:
    groups: list[tuple[str, str | None, list[Piece]]] = []
    colophon: Piece | None = None
    for piece in pieces:
        if piece.kind in {"title", "toc", "part"}:
            continue
        if piece.kind == "colophon":
            colophon = piece
            continue
        label = piece.toc_group or "The book"
        href = f"#{piece.group_anchor}" if piece.group_anchor else None
        if not groups or groups[-1][0] != label:
            groups.append((label, href, [piece]))
        else:
            groups[-1][2].append(piece)
    items = [
        '<?xml version="1.0" encoding="utf-8"?>',
        '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="en" xml:lang="en">',
        "<head>",
        "<title>Contents</title>",
        '<link rel="stylesheet" type="text/css" href="styles/epub.css"/>',
        "</head>",
        "<body>",
        '<nav epub:type="toc" id="toc">',
        "<h1>Contents</h1>",
        "<ol>",
        '<li><a href="text/title.xhtml">The Closer Project</a></li>',
    ]
    for label, href, members in groups:
        if href and href.startswith("#"):
            items.append(f'<li><a href="text/{esc(href[1:])}.xhtml">{esc(label)}</a>')
        else:
            items.append(f"<li><span>{esc(label)}</span>")
        items.append("<ol>")
        for piece in members:
            items.append(f'<li><a href="text/{esc(piece.anchor)}.xhtml">{esc(piece.title)}</a></li>')
        items.append("</ol></li>")
    if colophon is not None:
        items.append(f'<li><a href="text/{esc(colophon.anchor)}.xhtml">{esc(colophon.title)}</a></li>')
    items.extend(["</ol>", "</nav>", "</body>", "</html>"])
    return "\n".join(items) + "\n"


def rewrite_epub_images(inner: str, images: dict[str, Path]) -> str:
    def replace(match: re.Match[str]) -> str:
        rel = match.group(1)
        arc = "images/" + rel.replace("/", "-")
        images[arc] = ROOT / rel
        return f'src="../{arc}"'

    return re.sub(r'src="asset:([^"]+)"', replace, inner)


def file_for_anchor(pieces: list[Piece]) -> dict[str, str]:
    mapping: dict[str, str] = {"contents": "../nav.xhtml", "toc": "../nav.xhtml"}
    for piece in pieces:
        if piece.kind == "toc":
            continue
        filename = f"{piece.anchor}.xhtml"
        for found in re.findall(r'id="([^"]+)"', piece.inner):
            mapping[found] = filename
    return mapping


def rewrite_epub_links(inner: str, mapping: dict[str, str], current: str) -> str:
    def replace(match: re.Match[str]) -> str:
        anchor = match.group(1)
        target = mapping.get(anchor)
        if target is None or target == current:
            return f'href="#{anchor}"'
        if target.startswith("../"):
            return f'href="{target}#{anchor}"'
        return f'href="{target}#{anchor}"'

    return re.sub(r'href="#([^"]+)"', replace, inner)


def epub_package(pieces: list[Piece], css: str, proof: bool) -> bytes:
    mapping = file_for_anchor(pieces)
    images: dict[str, Path] = {}
    docs: list[tuple[str, str, str]] = []
    for piece in pieces:
        if piece.kind == "toc":
            continue
        inner = rewrite_epub_images(piece.inner, images)
        inner = rewrite_epub_links(inner, mapping, f"{piece.anchor}.xhtml")
        document = xhtml_document(piece.title, inner, "../styles/epub.css")
        docs.append((piece.anchor, f"text/{piece.anchor}.xhtml", document))
    nav = epub_nav(pieces)
    description = "Draft v1 digital reader. Printable writing lines live in the workbook; this EPUB is for reading."
    if proof:
        description += " Includes a typesetting fixture titled FIXTURE NOT A CHAPTER, which is not a catalog entry."
    manifest = [
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>',
        '<item id="style" href="styles/epub.css" media-type="text/css"/>',
    ]
    spine = ['<itemref idref="title-page"/>', '<itemref idref="nav"/>']
    for anchor, href, _document in docs:
        manifest.append(f'<item id="{esc(anchor)}" href="{esc(href)}" media-type="application/xhtml+xml"/>')
        if anchor != "title-page":
            spine.append(f'<itemref idref="{esc(anchor)}"/>')
    for arc in sorted(images):
        suffix = Path(arc).suffix.lower()
        media = {
            ".svg": "image/svg+xml",
            ".png": "image/png",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".gif": "image/gif",
            ".webp": "image/webp",
        }.get(suffix, "application/octet-stream")
        item_id = "img-" + slugify(Path(arc).stem)
        manifest.append(f'<item id="{esc(item_id)}" href="{esc(arc)}" media-type="{media}"/>')
    opf = f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book-id" xml:lang="en">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
<dc:identifier id="book-id">urn:closer-project:reader:v1-draft</dc:identifier>
<dc:title>The Closer Project</dc:title>
<dc:language>en</dc:language>
<dc:creator>The Closer Project</dc:creator>
<dc:description>{esc(description)}</dc:description>
<meta property="dcterms:modified">{STAMP}</meta>
</metadata>
<manifest>
{chr(10).join(manifest)}
</manifest>
<spine>
{chr(10).join(spine)}
</spine>
</package>
"""
    container = """<?xml version="1.0" encoding="utf-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles>
<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
</rootfiles>
</container>
"""
    files: list[tuple[str, bytes]] = [
        ("mimetype", b"application/epub+zip"),
        ("META-INF/container.xml", container.encode("utf-8")),
        ("OEBPS/content.opf", opf.encode("utf-8")),
        ("OEBPS/nav.xhtml", nav.encode("utf-8")),
        ("OEBPS/styles/epub.css", css.encode("utf-8")),
    ]
    for _anchor, href, document in docs:
        files.append((f"OEBPS/{href}", document.encode("utf-8")))
    for arc in sorted(images):
        files.append((f"OEBPS/{arc}", images[arc].read_bytes()))
    for name, data in files:
        if name.endswith((".xhtml", ".opf", ".xml")) and name != "mimetype":
            try:
                ET.fromstring(data)
            except ET.ParseError as exc:
                raise SystemExit(f"{name} is not well-formed XML: {exc}") from exc
    return zip_epub(files)


def zip_epub(files: list[tuple[str, bytes]]) -> bytes:
    from io import BytesIO

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for index, (name, data) in enumerate(files):
            info = zipfile.ZipInfo(filename=name, date_time=(2026, 10, 14, 0, 0, 0))
            info.compress_type = zipfile.ZIP_STORED if index == 0 else zipfile.ZIP_DEFLATED
            info.external_attr = 0
            info.create_system = 0
            info.flag_bits = 0
            info.extra = b""
            archive.writestr(info, data, compress_type=info.compress_type, compresslevel=None if index == 0 else 6)
    payload = buffer.getvalue()
    problems = epub_mimetype_problems(payload)
    if problems:
        raise SystemExit("EPUB mimetype entry is invalid: " + "; ".join(problems))
    return payload


def epub_mimetype_problems(payload: bytes) -> list[str]:
    problems = []
    if payload[:4] != b"PK\x03\x04":
        return ["missing zip header"]
    method = int.from_bytes(payload[8:10], "little")
    name_len = int.from_bytes(payload[26:28], "little")
    extra_len = int.from_bytes(payload[28:30], "little")
    comp_size = int.from_bytes(payload[18:22], "little")
    name = payload[30 : 30 + name_len]
    start = 30 + name_len + extra_len
    data = payload[start : start + comp_size]
    if name != b"mimetype":
        problems.append(f"first entry is {name!r}")
    if method != 0:
        problems.append(f"mimetype is compressed ({method})")
    if extra_len != 0:
        problems.append(f"mimetype has an extra field ({extra_len} bytes)")
    if data != b"application/epub+zip":
        problems.append(f"mimetype contents are {data!r}")
    return problems


def set_run_font(run, name: str, size: float, color: RGBColor, bold: bool = False, italic: bool = False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), name)


def set_style_font(style, name: str, size: float, color_hex: str, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(int(color_hex[0:2], 16), int(color_hex[2:4], 16), int(color_hex[4:6], 16))
    r_pr = style.element.get_or_add_rPr()
    for tag in ("w:rFonts", "w:color", "w:b"):
        for child in list(r_pr):
            if child.tag == qn(tag):
                r_pr.remove(child)
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), name)
    r_pr.append(r_fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), color_hex)
    r_pr.append(color)
    toggle = OxmlElement("w:b")
    toggle.set(qn("w:val"), "1" if bold else "0")
    r_pr.append(toggle)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    set_style_font(normal, "Georgia", 11, "241F1C")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, font, size, color, before, after in (
        ("Heading 1", "Palatino Linotype", 26, "241F1C", 0, 8),
        ("Heading 2", "Palatino Linotype", 16, "7C3A46", 16, 4),
        ("Heading 3", "Palatino Linotype", 13, "241F1C", 12, 2),
    ):
        style = document.styles[style_name]
        set_style_font(style, font, size, color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=140, right=140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def set_table_width(table, twips: int = CONTENT_TWIPS) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(twips))
    tbl_w.set(qn("w:type"), "dxa")
    table.autofit = False


def set_table_borders(table, color: str, sz: str = "8", inside: bool = False) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        draw = inside or not edge.startswith("inside")
        el.set(qn("w:val"), "single" if draw else "nil")
        el.set(qn("w:sz"), sz if draw else "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color if draw else "auto")
        borders.append(el)
    tbl_pr.append(borders)


def cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def box_cell(cell, color: str, sz: str = "8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def add_bottom_border(paragraph, color: str = "D9CFC0") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    set_run_font(run, "Palatino Linotype", 9, COPPER)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(XML_SPACE, "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def configure_section(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(7)
    section.page_height = Inches(10)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.32)
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("THE CLOSER PROJECT")
    set_run_font(run, "Palatino Linotype", 8, COPPER)
    section.first_page_header.paragraphs[0].text = ""
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer)
    section.first_page_footer.paragraphs[0].text = ""
    background = OxmlElement("w:background")
    background.set(qn("w:color"), "F6F1E7")
    document.element.insert(0, background)
    document.settings.element.append(OxmlElement("w:displayBackgroundShape"))
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    document.settings.element.append(update)


def add_internal_link(paragraph, bookmarks: Bookmarks, anchor: str, text: str, size: float = 11) -> None:
    href = normalize_href(anchor)
    target = href[1:] if href.startswith("#") else ""
    name = bookmarks.by_html.get(target)
    if not name:
        run = paragraph.add_run(text)
        set_run_font(run, "Georgia", size, INK)
        return
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), name)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), "Georgia")
    r_pr.append(r_fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "3E6464")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.set(XML_SPACE, "preserve")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_external_link(paragraph, url: str, text: str, size: float = 11) -> None:
    rel_id = paragraph.part.relate_to(url, HYPERLINK_REL, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "3E6464")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inlines(paragraph, text: str, bookmarks: Bookmarks, font: str, size: float, color: RGBColor, italic=False, bold=False) -> None:
    runs = parse_inline(text)
    if not runs:
        run = paragraph.add_run("")
        set_run_font(run, font, size, color, bold, italic)
        return
    for item in runs:
        kind = item[0]
        if kind == "link":
            href = normalize_href(item[2])
            if href.startswith("#"):
                add_internal_link(paragraph, bookmarks, href, item[1], size)
            elif href.startswith(("http://", "https://", "mailto:")):
                add_external_link(paragraph, href, item[1], size)
            else:
                run = paragraph.add_run(item[1])
                set_run_font(run, font, size, TEAL, bold, italic)
            continue
        if kind == "image":
            run = paragraph.add_run(item[1] or item[2])
            set_run_font(run, font, size, TEAL, italic=True)
            continue
        value = item[1]
        run = paragraph.add_run(value)
        if kind == "code":
            set_run_font(run, "Consolas", max(9, size - 1), color)
        else:
            set_run_font(
                run,
                font,
                size,
                color,
                bold or kind == "bold",
                italic or kind == "italic",
            )


def add_text_p(container, text, bookmarks: Bookmarks, *, style=None, font="Georgia", size=11, color=INK,
               italic=False, bold=False, align=None, before=0, after=8, keep=False, bookmark=None):
    paragraph = container.add_paragraph()
    if style:
        paragraph.style = style
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.keep_with_next = keep
    add_inlines(paragraph, text, bookmarks, font, size, color, italic, bold)
    if bookmark:
        bookmarks.add(paragraph, bookmark)
    return paragraph


def add_rule_line(container) -> None:
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = Pt(28)
    run = paragraph.add_run(" ")
    set_run_font(run, "Georgia", 11, RGBColor(0xFB, 0xF8, 0xF2))
    add_bottom_border(paragraph, "D9CFC0")


def fill_first(cell, builder) -> None:
    """Use the cell's starter paragraph, then extra paragraphs from the builder."""
    state = {"used": False, "cell": cell}

    def add():
        if not state["used"]:
            state["used"] = True
            return cell.paragraphs[0]
        return cell.add_paragraph()

    builder(add)


def add_journal_docx(document: Document, text: str, bookmarks: Bookmarks) -> None:
    sections = parse_journal(text)
    table = document.add_table(rows=1 + len(sections), cols=1)
    set_table_width(table)
    set_table_borders(table, "D9CFC0", inside=False)
    intro = table.cell(0, 0)
    shade_cell(intro, "FBF8F2")
    set_cell_margins(intro)
    cant_split(table.rows[0])

    def build_intro(add):
        kicker = add()
        kicker.paragraph_format.space_after = Pt(2)
        run = kicker.add_run("OUR STORY")
        set_run_font(run, "Calibri", 9, COPPER)
        note = add()
        note.paragraph_format.space_after = Pt(4)
        run = note.add_run("Sharing either private page is optional.")
        set_run_font(run, "Georgia", 10, INK, italic=True)

    fill_first(intro, build_intro)
    for index, section in enumerate(sections, start=1):
        cell = table.cell(index, 0)
        shade_cell(cell, "FBF8F2")
        set_cell_margins(cell)
        cant_split(table.rows[index])

        def build(add, section=section):
            heading = add()
            heading.style = document.styles["Heading 3"]
            heading.paragraph_format.space_before = Pt(6)
            heading.paragraph_format.space_after = Pt(2)
            run = heading.add_run(journal_label(section["key"]))
            set_run_font(run, "Palatino Linotype", 13, INK)
            if section["intro"]:
                instruction = add()
                instruction.paragraph_format.space_after = Pt(2)
                add_inlines(instruction, section["intro"], bookmarks, "Georgia", 10, INK, italic=True)
            if section["items"]:
                for item in section["items"]:
                    prompt = add()
                    prompt.paragraph_format.space_before = Pt(6)
                    prompt.paragraph_format.space_after = Pt(1)
                    prompt.paragraph_format.keep_with_next = True
                    add_inlines(prompt, item, bookmarks, "Georgia", 11, INK)
                    for _ in range(8):
                        add_rule_line(cell)
            else:
                for _ in range(2 if section["key"] == "session" else 8):
                    add_rule_line(cell)

        fill_first(cell, build)
    document.add_paragraph().paragraph_format.space_after = Pt(6)


def add_cards_docx(document: Document, text: str, bookmarks: Bookmarks) -> None:
    cards = parse_cards(text)
    if not cards:
        return
    columns = 2
    row_count = (len(cards) + columns - 1) // columns
    table = document.add_table(rows=row_count, cols=columns)
    set_table_width(table)
    set_table_borders(table, "F6F1E7", inside=False)
    for index, (title, body) in enumerate(cards):
        cell = table.cell(index // columns, index % columns)
        shade_cell(cell, "FBF8F2")
        box_cell(cell, "A85A32", "8")
        set_cell_margins(cell, 100, 80, 120, 120)

        def build(add, title=title, body=body):
            if title:
                heading = add()
                heading.paragraph_format.space_after = Pt(2)
                run = heading.add_run(title)
                set_run_font(run, "Palatino Linotype", 13, BURGUNDY)
            if body:
                paragraph = add()
                paragraph.paragraph_format.space_after = Pt(2)
                add_inlines(paragraph, body, bookmarks, "Georgia", 10.5, INK)

        fill_first(cell, build)
    if len(cards) % columns == 1:
        last = table.rows[-1]
        last.cells[0].merge(last.cells[1])
    document.add_paragraph().paragraph_format.space_after = Pt(6)


def add_table_docx(document: Document, header: list[str], rows: list[list[str]], bookmarks: Bookmarks) -> None:
    table = document.add_table(rows=1 + len(rows), cols=len(header))
    set_table_width(table)
    set_table_borders(table, "D9CFC0", sz="4", inside=True)
    for index, text in enumerate(header):
        cell = table.cell(0, index)
        shade_cell(cell, "F6F1E7")
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(2)
        add_inlines(paragraph, text, bookmarks, "Palatino Linotype", 11, INK)
        add_bottom_border(paragraph, "A85A32")
    for r_index, row in enumerate(rows):
        padded = list(row) + [""] * max(0, len(header) - len(row))
        for c_index, text in enumerate(padded[: len(header)]):
            cell = table.cell(r_index + 1, c_index)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            add_inlines(paragraph, text, bookmarks, "Georgia", 10, INK)
    document.add_paragraph().paragraph_format.space_after = Pt(6)


def add_blocks_docx(document: Document, doc: SourceDoc, bookmarks: Bookmarks) -> None:
    visible = body_blocks(doc)
    close_next = False
    for block, heading_id in zip(visible, doc.heading_ids):
        kind = block[0]
        if kind == "heading":
            if cover_heading(block[2]) or close_next:
                document.add_page_break()
            close_next = closes_player_text(block[2])
            level = {1: "Heading 1", 2: "Heading 2"}.get(block[1], "Heading 3")
            font = "Palatino Linotype"
            size = {1: 26, 2: 16}.get(block[1], 13)
            color = BURGUNDY if block[1] == 2 else INK
            add_text_p(
                document,
                block[2],
                bookmarks,
                style=level,
                font=font,
                size=size,
                color=color,
                before=14 if block[1] != 1 else 0,
                after=4,
                keep=True,
                bookmark=heading_id,
            )
        elif kind == "para":
            parsed = parse_inline(block[1])
            if len(parsed) == 1 and parsed[0][0] == "image":
                add_docx_image(document, parsed[0][2], parsed[0][1] or "Illustration")
            else:
                add_text_p(document, block[1], bookmarks)
        elif kind == "list":
            style = "List Number" if block[1] else "List Bullet"
            for item in block[2]:
                add_text_p(document, item, bookmarks, style=style, after=3)
        elif kind == "table":
            add_table_docx(document, block[1], block[2], bookmarks)
        elif kind == "fence" and block[1] == "journal":
            add_journal_docx(document, block[2], bookmarks)
        elif kind == "fence" and block[1] == "cards":
            add_cards_docx(document, block[2], bookmarks)
        elif kind == "fence":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block[2])
            set_run_font(run, "Consolas", 9, INK)
        elif kind == "hr":
            paragraph = document.add_paragraph()
            add_bottom_border(paragraph, "A85A32")


def svg_pixel_size(path: Path) -> tuple[int, int]:
    head = path.read_text(encoding="utf-8")[:800]
    match = re.search(r'width="(\d+(?:\.\d+)?)"\s+height="(\d+(?:\.\d+)?)"', head)
    if not match:
        return 440, 280
    return int(float(match.group(1))), int(float(match.group(2)))


def raster_svg(path: Path) -> Path | None:
    """Cache a parchment PNG of an SVG so the DOCX can embed a drawing."""
    digest = hashlib.sha256(path.read_bytes()).hexdigest()[:16]
    out = ROOT / "build" / "raster" / f"{digest}.png"
    if out.is_file() and out.stat().st_size > 800:
        return out
    out.parent.mkdir(parents=True, exist_ok=True)
    browsers = [item for item in BROWSERS if item.is_file()]
    if not browsers:
        return None
    width, height = svg_pixel_size(path)
    with tempfile.TemporaryDirectory(prefix="closer-raster-", ignore_cleanup_errors=True) as profile:
        command = [
            str(browsers[0]),
            "--headless=new",
            "--disable-gpu",
            "--hide-scrollbars",
            "--no-first-run",
            "--no-default-browser-check",
            "--disable-extensions",
            "--allow-file-access-from-files",
            "--default-background-color=FFF6F1E7",
            f"--window-size={max(width, 1)},{max(height, 1)}",
            f"--user-data-dir={profile}",
            f"--screenshot={out.resolve()}",
            path.resolve().as_uri(),
        ]
        try:
            subprocess.run(
                command,
                cwd=str(ROOT),
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=45,
                check=False,
            )
        except (subprocess.TimeoutExpired, OSError):
            return None
    if out.is_file() and out.stat().st_size > 800:
        return out
    return None


def picture_width(path: Path):
    if path.suffix.lower() == ".svg":
        width, height = svg_pixel_size(path)
    else:
        width, height = 440, 280
    width_in = width / 100
    height_in = height / 100
    scale = min(1.0, 5.2 / width_in if width_in else 1, 6.2 / height_in if height_in else 1)
    return Inches(max(1.2, width_in * scale))


def prefetch_rasters(paths: list[Path]) -> None:
    unique: list[Path] = []
    seen: set[str] = set()
    for path in paths:
        key = str(path.resolve())
        if key not in seen:
            seen.add(key)
            unique.append(path)
    pending = []
    for path in unique:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()[:16]
        cached = ROOT / "build" / "raster" / f"{digest}.png"
        if not (cached.is_file() and cached.stat().st_size > 800):
            pending.append(path)
    if not pending:
        return
    from concurrent.futures import ThreadPoolExecutor

    with ThreadPoolExecutor(max_workers=4) as pool:
        list(pool.map(raster_svg, pending))


def workbook_svgs(docs: list[SourceDoc]) -> list[Path]:
    found: list[Path] = []
    for rel in (
        "assets/covers/cover-front.svg",
        "assets/covers/cover-back.svg",
        *(f"assets/parts/{index:02d}-opener.svg" for index in range(1, 14)),
        *(f"assets/signatures/S{index:02d}-hero.svg" for index in range(1, 7)),
    ):
        path = resolve_asset(rel)
        if path is not None:
            found.append(path)
    for doc in docs:
        motif = doc.meta.get("motif")
        if not motif:
            continue
        path = resolve_asset(str(motif))
        if path is not None and path.suffix.lower() == ".svg":
            found.append(path)
    return found


def add_docx_image(document: Document, src: str, alt: str, missing_note: bool = False, svg_label: str = "Motif") -> None:
    path = resolve_asset(src)
    if path is None:
        if missing_note:
            add_note = document.add_paragraph()
            run = add_note.add_run(f"No image placed. The file is not on disk: {src}")
            set_run_font(run, "Calibri", 9, TEAL, italic=True)
        return
    draw_from = path
    if path.suffix.lower() == ".svg":
        png = raster_svg(path)
        if png is None:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(f"{svg_label} (drawing is in the HTML, EPUB, and PDF): {rel_posix(path)}. {alt}")
            set_run_font(run, "Calibri", 9, TEAL, italic=True)
            return
        draw_from = png
    try:
        from PIL import Image

        with Image.open(draw_from) as image:
            image.verify()
        document.add_picture(str(draw_from), width=picture_width(path))
    except Exception as exc:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"Image could not be embedded ({path.name}): {exc}")
        set_run_font(run, "Calibri", 9, TEAL, italic=True)


def add_opener_docx(document: Document, doc: SourceDoc, bookmarks: Bookmarks) -> None:
    if doc.kind == "fixture":
        banner = add_text_p(document, "FIXTURE NOT A CHAPTER. This is a typesetting sample, not a catalog entry.",
                            bookmarks, font="Calibri", size=10, before=0, after=8)
        shade_paragraph(banner, "FBF8F2")
        add_text_p(document, "FIXTURE · NOT A CATALOG ENTRY", bookmarks, font="Calibri", size=9, color=COPPER, after=0)
        add_text_p(document, "PROOF", bookmarks, font="Palatino Linotype", size=11, color=COPPER, after=2)
    elif doc.kind == "signature":
        add_text_p(document, "SIGNATURE SESSION", bookmarks, font="Calibri", size=9, color=COPPER, after=0)
    else:
        part = doc.meta.get("part")
        part_title = str(doc.meta.get("part_title") or "")
        if part and part_title:
            kicker = f"PART {part} · {part_title}"
        elif part_title:
            kicker = part_title
        elif part:
            kicker = f"PART {part}"
        else:
            kicker = ""
        if kicker:
            add_text_p(document, kicker.upper(), bookmarks, font="Calibri", size=9, color=COPPER, after=0)
        if doc.meta.get("id") is not None:
            raw = str(doc.meta["id"])
            shown = raw.zfill(3) if raw.isdigit() else raw
            add_text_p(document, shown, bookmarks, font="Palatino Linotype", size=11, color=COPPER, after=2)
    add_text_p(document, doc.title, bookmarks, style="Heading 1", font="Palatino Linotype", size=26, after=6,
               keep=True, bookmark=doc.anchor)
    if doc.meta.get("deck"):
        add_text_p(document, str(doc.meta["deck"]), bookmarks, size=12, italic=True, after=6)
    if doc.meta.get("motif"):
        add_docx_image(document, str(doc.meta["motif"]), f"Motif for {doc.title}", missing_note=doc.kind == "fixture")
    summary = meta_line(doc)
    if summary:
        add_text_p(document, summary, bookmarks, font="Calibri", size=9, color=TEAL, after=4)
    links = doc.meta.get("cross_links") or []
    if isinstance(links, list) and links:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        lead = paragraph.add_run("See also ")
        set_run_font(lead, "Calibri", 9, TEAL)
        for index, item in enumerate(links):
            if index:
                comma = paragraph.add_run(", ")
                set_run_font(comma, "Calibri", 9, TEAL)
            raw = str(item).strip()
            if raw.isdigit():
                add_internal_link(paragraph, bookmarks, f"#chapter-{int(raw):03d}", raw.zfill(3), 9)
            else:
                plain = paragraph.add_run(raw)
                set_run_font(plain, "Calibri", 9, TEAL)


def add_indexes_docx(document: Document, bookmarks: Bookmarks, docs: list[SourceDoc]) -> None:
    for section in build_index_sections(docs):
        document.add_page_break()
        add_text_p(document, "INDEX", bookmarks, font="Calibri", size=9, color=COPPER, after=2)
        add_text_p(
            document,
            section.title,
            bookmarks,
            style="Heading 1",
            font="Palatino Linotype",
            size=26,
            after=4,
            bookmark=section.anchor,
        )
        add_text_p(
            document,
            "Generated from the chapter frontmatter at build time.",
            bookmarks,
            font="Calibri",
            size=9,
            color=TEAL,
            after=6,
        )
        for block in section.blocks:
            if block[0] == "note":
                add_text_p(document, block[1], bookmarks)
            elif block[0] == "rows":
                for ref in block[1]:
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.space_before = Pt(1)
                    paragraph.paragraph_format.space_after = Pt(1)
                    add_internal_link(paragraph, bookmarks, f"#{ref['anchor']}", ref["title"], 11)
                    ident = paragraph.add_run("  " + ref["shown"])
                    set_run_font(ident, "Calibri", 10, TEAL)
                    add_bottom_border(paragraph, "C6A15B")
            elif block[0] == "group":
                add_text_p(
                    document,
                    block[1],
                    bookmarks,
                    style="Heading 2",
                    font="Palatino Linotype",
                    size=16,
                    color=BURGUNDY,
                    before=12,
                    after=4,
                    keep=True,
                )
                for ref in block[2]:
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(1)
                    label = f"{ref['shown']}  {ref['title']}".strip()
                    add_internal_link(paragraph, bookmarks, f"#{ref['anchor']}", label, 11)


def register_bookmarks(docs: list[SourceDoc], part_anchors: dict[str, str]) -> Bookmarks:
    bookmarks = Bookmarks()
    for html_id in ("title-page", "contents", "colophon"):
        bookmarks.register(html_id)
    for anchor, _title in INDEX_SPECS:
        bookmarks.register(anchor)
    for anchor in part_anchors.values():
        bookmarks.register(anchor)
    for doc in docs:
        bookmarks.register(doc.anchor)
        for heading_id in doc.heading_ids:
            if heading_id:
                bookmarks.register(heading_id)
    return bookmarks


def set_core_properties(document: Document, proof: bool) -> None:
    props = document.core_properties
    props.title = "The Closer Project — Illustrated Workbook"
    props.subject = "Draft v1 typesetting proof" if proof else "Draft v1"
    props.author = "The Closer Project"
    props.category = "Draft"
    props.comments = "Generated from Markdown by scripts/build_all.py. Not a release."
    props.last_modified_by = "The Closer Project"


def build_docx(docs: list[SourceDoc], part_anchors: dict[str, str], proof: bool, bookmarks: Bookmarks) -> None:
    document = Document()
    configure_styles(document)
    configure_section(document)
    set_core_properties(document, proof)
    prefetch_rasters(workbook_svgs(docs))
    add_text_p(document, "DRAFT V1 · 14 OCTOBER 2026", bookmarks, font="Calibri", size=9, color=COPPER,
               align=WD_ALIGN_PARAGRAPH.CENTER, before=72, after=6, bookmark="title-page")
    add_docx_image(
        document,
        "assets/covers/cover-front.svg",
        "Front cover of The Closer Project.",
        missing_note=False,
        svg_label="Cover",
    )
    add_text_p(document, "The Closer Project", bookmarks, style="Heading 1", font="Palatino Linotype", size=32,
               align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_text_p(document, "ILLUSTRATED WORKBOOK", bookmarks, font="Palatino Linotype", size=14, color=COPPER,
               align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_text_p(document, "Experiences for two adults who already like each other.", bookmarks, size=12, italic=True,
               align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    if proof:
        banner = add_text_p(
            document,
            "Typesetting proof. The fixture at the end is not a chapter and is not catalog entry 001.",
            bookmarks,
            font="Calibri",
            size=10,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        shade_paragraph(banner, "FBF8F2")
    document.add_page_break()
    add_text_p(document, "Contents", bookmarks, style="Heading 1", font="Palatino Linotype", size=26, bookmark="contents")
    current = None
    for doc in docs:
        if doc.kind == "front":
            label, href = "Front matter", None
        elif doc.kind == "chapter":
            part = doc.meta.get("part")
            part_title = str(doc.meta.get("part_title") or "")
            label = f"Part {part}" + (f" · {part_title}" if part_title else "") if part else "Chapters"
            key = str(part) if part is not None else ""
            href = part_anchors.get(key)
        elif doc.kind == "signature":
            label, href = "Signature sessions", None
        elif doc.kind == "bibliography":
            continue
        else:
            label, href = "Typesetting fixture (not a chapter)", None
        if label != current:
            current = label
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(2)
            if href:
                add_internal_link(paragraph, bookmarks, f"#{href}", label, 14)
            else:
                run = paragraph.add_run(label)
                set_run_font(run, "Palatino Linotype", 14, BURGUNDY)
        entry = document.add_paragraph()
        entry.paragraph_format.space_before = Pt(1)
        entry.paragraph_format.space_after = Pt(1)
        add_internal_link(entry, bookmarks, f"#{doc.anchor}", doc.title, 11)
        add_bottom_border(entry, "C6A15B")
    index_label = document.add_paragraph()
    index_label.paragraph_format.space_before = Pt(12)
    index_label.paragraph_format.space_after = Pt(2)
    index_run = index_label.add_run("Indexes")
    set_run_font(index_run, "Palatino Linotype", 14, BURGUNDY)
    for anchor, title in INDEX_SPECS:
        entry = document.add_paragraph()
        entry.paragraph_format.space_before = Pt(1)
        entry.paragraph_format.space_after = Pt(1)
        add_internal_link(entry, bookmarks, f"#{anchor}", title, 11)
        add_bottom_border(entry, "C6A15B")
    for doc in docs:
        if doc.kind != "bibliography":
            continue
        entry = document.add_paragraph()
        entry.paragraph_format.space_before = Pt(10)
        entry.paragraph_format.space_after = Pt(1)
        add_internal_link(entry, bookmarks, f"#{doc.anchor}", doc.title, 14)
        add_bottom_border(entry, "C6A15B")
    document.add_page_break()
    seen_parts: set[str] = set()
    at_break = True
    for doc in docs:
        if doc.kind == "bibliography":
            continue
        if doc.kind == "chapter":
            key = str(doc.meta.get("part")) if doc.meta.get("part") is not None else ""
            if key and key not in seen_parts and key in part_anchors:
                seen_parts.add(key)
                if not at_break:
                    document.add_page_break()
                part_title = str(doc.meta.get("part_title") or f"Part {key}")
                add_text_p(document, f"PART {key}", bookmarks, font="Calibri", size=10, color=COPPER, before=28, after=4)
                opener_src = part_opener_path(key)
                if opener_src:
                    add_docx_image(document, opener_src, part_title, missing_note=False, svg_label="Part opener")
                add_text_p(document, part_title, bookmarks, style="Heading 1", font="Palatino Linotype", size=28,
                           before=0, after=8, bookmark=part_anchors[key])
                intro = part_introductions().get(str(int(key)) if key.isdigit() else "", "")
                if intro:
                    add_text_p(document, intro, bookmarks, size=12, after=8)
                document.add_page_break()
                at_break = True
        if doc.kind != "front" and not at_break:
            document.add_page_break()
            at_break = True
        if doc.kind in {"chapter", "signature", "fixture"}:
            add_opener_docx(document, doc, bookmarks)
        else:
            # Front matter: bookmark the first heading if present, and the article anchor on that heading.
            pass
        if doc.kind == "front":
            add_front_docx(document, doc, bookmarks)
        else:
            add_blocks_docx(document, doc, bookmarks)
        at_break = False
    add_indexes_docx(document, bookmarks, docs)
    for doc in docs:
        if doc.kind != "bibliography":
            continue
        document.add_page_break()
        add_front_docx(document, doc, bookmarks)
    document.add_page_break()
    add_docx_image(
        document,
        "assets/covers/cover-back.svg",
        "Back cover of The Closer Project.",
        missing_note=False,
        svg_label="Back cover",
    )
    add_text_p(document, "Build note", bookmarks, style="Heading 1", font="Palatino Linotype", size=20, bookmark="colophon")
    add_text_p(
        document,
        "Set 14 October 2026 from Markdown. This file is a draft. The private journal folder is not read.",
        bookmarks,
    )
    add_text_p(document, "Sources in this file:", bookmarks, keep=True)
    for doc in docs:
        add_text_p(document, doc.rel, bookmarks, font="Consolas", size=9, after=2)
    tmp = DOCX_PATH.with_name(DOCX_PATH.name + ".tmp")
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(tmp)
    normalize_office_zip(tmp, DOCX_PATH)


def add_front_docx(document: Document, doc: SourceDoc, bookmarks: Bookmarks) -> None:
    visible = body_blocks(doc)
    marked = False
    ids = list(doc.heading_ids)
    for block, heading_id in zip(visible, ids):
        if block[0] == "heading" and block[1] == 1 and not marked:
            add_text_p(
                document,
                block[2],
                bookmarks,
                style="Heading 1",
                font="Palatino Linotype",
                size=26,
                before=16,
                after=8,
                keep=True,
                bookmark=doc.anchor,
            )
            if heading_id and heading_id != doc.anchor:
                bookmarks.add(document.paragraphs[-1], heading_id)
            marked = True
            continue
        # Reuse the block walker for the remainder by a one-off wrapper.
        shadow = SourceDoc(
            path=doc.path,
            rel=doc.rel,
            kind="front",
            meta=doc.meta,
            body=doc.body,
            blocks=[block] if block[0] != "heading" or block[1] != 1 else [block],
            anchor=doc.anchor,
            title=doc.title,
            heading_ids=[heading_id],
        )
        if block[0] == "heading" and block[1] == 1:
            continue
        add_blocks_docx(document, shadow, bookmarks)
    if not marked:
        add_text_p(document, doc.title, bookmarks, style="Heading 1", bookmark=doc.anchor)


def normalize_office_zip(src: Path, dest: Path) -> None:
    mapping: dict[str, str] = {}

    def freeze(xml: str) -> str:
        def repl(match: re.Match[str]) -> str:
            value = match.group(2)
            if value not in mapping:
                mapping[value] = f"{len(mapping) + 1:08X}"
            return f'{match.group(1)}{mapping[value]}"'

        xml = re.sub(r'(w:rsid[A-Za-z]*=")([0-9A-Fa-f]+)"', repl, xml)
        xml = re.sub(r'(w14:paraId=")([0-9A-Fa-f]+)"', repl, xml)
        xml = re.sub(r'(w14:textId=")([0-9A-Fa-f]+)"', repl, xml)
        xml = re.sub(
            r"<dcterms:created\b[^>]*>.*?</dcterms:created>",
            '<dcterms:created xsi:type="dcterms:W3CDTF">2026-10-14T00:00:00Z</dcterms:created>',
            xml,
            flags=re.DOTALL,
        )
        xml = re.sub(
            r"<dcterms:modified\b[^>]*>.*?</dcterms:modified>",
            '<dcterms:modified xsi:type="dcterms:W3CDTF">2026-10-14T00:00:00Z</dcterms:modified>',
            xml,
            flags=re.DOTALL,
        )
        return xml

    with zipfile.ZipFile(src, "r") as archive:
        items = [(info.filename, archive.read(info.filename)) for info in archive.infolist()]
    written = dest.with_name(dest.name + ".norm")
    with zipfile.ZipFile(written, "w") as archive:
        for name, data in items:
            if name.endswith((".xml", ".rels")):
                data = freeze(data.decode("utf-8")).encode("utf-8")
            info = zipfile.ZipInfo(filename=name, date_time=(2026, 10, 14, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0
            info.create_system = 0
            archive.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED, compresslevel=6)
    src.unlink(missing_ok=True)
    written.replace(dest)


def print_pdf(
    html_path: Path,
    pdf_path: Path,
    error_path: Path | None = None,
    timeout: int = 1200,
    kept: str = "HTML, DOCX, and EPUB were kept.",
) -> None:
    if error_path is None:
        error_path = PDF_ERROR_PATH
    if pdf_path.exists():
        pdf_path.unlink()
    errors = []
    browsers = [path for path in BROWSERS if path.is_file()]
    if not browsers:
        message = f"PDF print failed. Neither Chrome nor Edge was found. {kept}"
        error_path.parent.mkdir(parents=True, exist_ok=True)
        error_path.write_text(message + "\n", encoding="utf-8")
        raise SystemExit(message + f" Exact error written to {error_path}.")
    uri = html_path.resolve().as_uri()
    for browser in browsers:
        for headless in ("--headless=new", "--headless"):
            with tempfile.TemporaryDirectory(prefix="closer-print-", ignore_cleanup_errors=True) as profile:
                if pdf_path.exists():
                    pdf_path.unlink()
                command = [
                    str(browser),
                    headless,
                    "--disable-gpu",
                    "--no-first-run",
                    "--no-default-browser-check",
                    "--disable-extensions",
                    "--allow-file-access-from-files",
                    "--no-pdf-header-footer",
                    "--hide-scrollbars",
                    f"--user-data-dir={profile}",
                    f"--print-to-pdf={pdf_path.resolve()}",
                    uri,
                ]
                try:
                    proc = subprocess.run(
                        command,
                        cwd=str(ROOT),
                        capture_output=True,
                        text=True,
                        encoding="utf-8",
                        errors="replace",
                        timeout=timeout,
                        check=False,
                    )
                except subprocess.TimeoutExpired as exc:
                    errors.append(f"{browser.name} {headless} timed out\n{exc}")
                    continue
                valid = pdf_path.is_file() and pdf_path.stat().st_size > 2000 and pdf_path.read_bytes()[:5] == b"%PDF-"
                if valid:
                    print(f"PASS  pdf via {browser.name} {headless} ({pdf_path.stat().st_size} bytes)")
                    return
                errors.append(
                    f"{browser.name} {headless} exit {proc.returncode}\nSTDOUT:\n{proc.stdout}\nSTDERR:\n{proc.stderr}"
                )
    error_path.parent.mkdir(parents=True, exist_ok=True)
    error_path.write_text("\n\n".join(errors), encoding="utf-8")
    raise SystemExit(f"PDF print failed. {kept} Exact error written to {error_path}.")


def assert_proof_html(page: str) -> None:
    required = [
        "FIXTURE NOT A CHAPTER",
        'id="contents"',
        'class="journal-spread"',
        'class="card-grid"',
        'class="writing-lines"',
        "<table>",
        'href="#fixture-not-a-chapter"',
    ]
    missing = [token for token in required if token not in page]
    if missing:
        raise SystemExit("Proof HTML is missing: " + ", ".join(missing))


def self_check() -> None:
    sample = "\n".join(f"## {heading}" for heading in REQUIRED_HEADINGS)
    if heading_problems(sample):
        raise SystemExit("Heading checker failed its own sample.")
    if "Why it might matter" not in heading_problems("## The hook\n## Notes\n"):
        raise SystemExit("Heading checker did not catch a missing heading.")
    journal = parse_journal("session: date, place\npartner_a:\n- one\n- two\n")
    if journal[0]["key"] != "session" or journal[1]["items"] != ["one", "two"]:
        raise SystemExit("Journal parser failed its own sample.")
    blocks = parse_blocks("| A | B |\n| --- | --- |\n| 1 | 2 |\n\n```cards\nAsk | Once\n```\n", "self-check")
    if blocks[0][0] != "table" or blocks[1] != ("fence", "cards", "Ask | Once"):
        raise SystemExit(f"Block parser failed its own sample: {blocks}")
    bold = parse_inline("**Established.** A practice")
    if bold[0] != ("bold", "Established."):
        raise SystemExit(f"Inline parser failed its own sample: {bold}")
    try:
        assert_readable(ROOT / "journal" / "private" / "unused.md")
    except SystemExit as exc:
        if "Refusing" not in str(exc):
            raise
    else:
        raise SystemExit("Private-path guard did not refuse journal/private.")
    signature_anchor = make_anchor("signature", {"id": "S01"}, "signature/s01.md", "self-check")
    if signature_anchor.startswith("chapter-") or signature_anchor != "signature-s01":
        raise SystemExit(f"Signature id S01 must not use the numeric chapter anchor: {signature_anchor}")
    incomplete = SourceDoc(
        path=ROOT / "signature" / "s01.md",
        rel="signature/s01.md",
        kind="signature",
        meta={"id": "S01"},
        body="## The hook\n## Notes\n",
        blocks=[],
        anchor=signature_anchor,
        title="Sample",
    )
    problems = signature_problems(incomplete)
    if not any("motif" in item for item in problems) or not any("Why it might matter" in item for item in problems):
        raise SystemExit(f"Signature check missed a heading or motif path: {problems}")
    complete = SourceDoc(
        path=incomplete.path,
        rel=incomplete.rel,
        kind="signature",
        meta={"id": "S01", "motif": "assets/signatures/s01/motif.svg"},
        body="\n".join(f"## {heading}" for heading in REQUIRED_HEADINGS),
        blocks=[],
        anchor=signature_anchor,
        title="Sample",
    )
    if signature_problems(complete):
        raise SystemExit("A complete signature was rejected.")
    check_signatures([])
    if part_opener_path(1) != "assets/parts/01-opener.svg" or part_opener_path(13) != "assets/parts/13-opener.svg":
        raise SystemExit("Part opener paths must use a zero-padded part number.")
    sink = ImageSink()
    opened = part_inner("1", "The Art of Really Knowing Someone", "part-01", sink)
    if "assets/parts/01-opener.svg" not in opened or "<figure" not in opened:
        raise SystemExit("Part opener did not include the SVG.")
    skipped = part_inner("99", "Missing part", "part-99", sink)
    if "<figure" in skipped:
        raise SystemExit("A missing part opener SVG should be skipped.")
    sample_docs = [
        SourceDoc(
            path=ROOT / "chapters" / "part-01" / "002.md",
            rel="chapters/part-01/002.md",
            kind="chapter",
            meta={
                "id": "002",
                "time_band": "5-10",
                "emotional_intensity": "light",
                "energy": "low",
                "formats": ["game"],
                "moods": ["laugh"],
                "repeatable": False,
            },
            body="",
            blocks=[],
            anchor="chapter-002",
            title="Beta",
        ),
        SourceDoc(
            path=ROOT / "chapters" / "part-01" / "001.md",
            rel="chapters/part-01/001.md",
            kind="chapter",
            meta={
                "id": "001",
                "time_band": "20-45",
                "emotional_intensity": "deep",
                "energy": "high",
                "formats": ["conversation"],
                "moods": ["discover"],
                "repeatable": True,
            },
            body="",
            blocks=[],
            anchor="chapter-001",
            title="Alpha",
        ),
        complete,
    ]
    sections = build_index_sections(sample_docs)
    by_anchor = {section.anchor: section for section in sections}
    alpha_rows = by_anchor["index-alphabetical"].blocks[0][1]
    if [ref["title"] for ref in alpha_rows] != ["Alpha", "Beta"]:
        raise SystemExit(f"Alphabetical index order failed: {alpha_rows}")
    if any(ref["anchor"].startswith("signature") for ref in alpha_rows):
        raise SystemExit("Signature sessions were forced into the chapter index.")
    time_labels = [block[1] for block in by_anchor["index-time"].blocks if block[0] == "group"]
    if time_labels[:2] != ["5-10", "20-45"]:
        raise SystemExit(f"Time index order failed: {time_labels}")
    revisit_rows = by_anchor["index-revisit"].blocks[0][1]
    if [ref["anchor"] for ref in revisit_rows] != ["chapter-001"]:
        raise SystemExit("Revisit index should list only chapters marked repeatable.")


def write_bytes(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_bytes(data)
    tmp.replace(path)


def main(argv: list[str] | None = None) -> int:
    configure_stdio()
    parser = argparse.ArgumentParser(description="Build The Closer Project draft workbook and reader.")
    parser.add_argument(
        "--proof",
        action="store_true",
        help="Include build/fixtures/typesetting-proof.md. The fixture is not a chapter.",
    )
    args = parser.parse_args(argv)
    self_check()
    assert_sources_unchanged()
    if not WORKBOOK_CSS.is_file() or not EPUB_CSS.is_file():
        raise SystemExit("Missing styles/workbook.css or styles/epub.css. No files were written.")
    docs, part_anchors = load_sources(args.proof)
    images = ImageSink()
    pieces = build_pieces(docs, part_anchors, args.proof, images)
    css = WORKBOOK_CSS.read_text(encoding="utf-8")
    epub_css = EPUB_CSS.read_text(encoding="utf-8")
    page = html_document(pieces, css)
    if args.proof:
        assert_proof_html(page)
    if "journal/private/" in page and "is not read" not in page:
        raise SystemExit("Build output mentioned the private journal without the refusal. No files were written.")
    write_bytes(HTML_PATH, page.encode("utf-8"))
    bookmarks = register_bookmarks(docs, part_anchors)
    build_docx(docs, part_anchors, args.proof, bookmarks)
    write_bytes(EPUB_PATH, epub_package(pieces, epub_css, args.proof))
    for doc in docs:
        label = "fixture, not a chapter" if doc.kind == "fixture" else doc.kind
        print(f"INCLUDED  {doc.rel}  ({label})")
    for note in images.notes:
        print(note)
    try:
        print_pdf(HTML_PATH, PDF_PATH)
    except SystemExit:
        print(f"KEPT  {HTML_PATH} ({HTML_PATH.stat().st_size} bytes)")
        print(f"KEPT  {DOCX_PATH} ({DOCX_PATH.stat().st_size} bytes)")
        print(f"KEPT  {EPUB_PATH} ({EPUB_PATH.stat().st_size} bytes)")
        raise
    if PDF_ERROR_PATH.exists():
        PDF_ERROR_PATH.unlink()
    for path in (HTML_PATH, DOCX_PATH, EPUB_PATH, PDF_PATH):
        print(f"WROTE  {path.relative_to(ROOT).as_posix()}  {path.stat().st_size} bytes")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception as exc:
        raise SystemExit(f"Build failed: {exc}") from exc
