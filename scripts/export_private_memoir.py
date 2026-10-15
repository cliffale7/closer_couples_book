#!/usr/bin/env python3
"""Build a local Our Story memoir from notes the couple typed themselves.

The tool reads only a private root you name. It writes only under that root's
exports folder. It does not open a network connection, and it does not print
the text of an entry. Tests should pass a temporary private root. Do not point
it at a real journal unless you mean to.

    python scripts/export_private_memoir.py --private-root journal/private --mode shared
    python scripts/export_private_memoir.py --private-root journal/private --mode partner-a
    python scripts/export_private_memoir.py --private-root journal/private --mode full-private --authorize-partner-a --authorize-partner-b
"""

from __future__ import annotations

import argparse
import html
import sys
from datetime import date
from pathlib import Path

import yaml
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

import build_all  # noqa: E402

MODES = ("shared", "partner-a", "partner-b", "full-private")
BODY_FILES = {"partner_a": "partner-a.md", "partner_b": "partner-b.md", "shared": "shared.md"}


def configure_stdio() -> None:
    build_all.configure_stdio()


def fail(message: str) -> None:
    raise SystemExit(message)


def is_inside(path: Path, root: Path) -> bool:
    resolved = path.resolve()
    base = root.resolve()
    return resolved == base or base in resolved.parents


def reject_symlink(path: Path) -> None:
    if path.is_symlink():
        fail(f"Refusing a symlink: {path}")


def assert_private_root(root: Path) -> Path:
    reject_symlink(root)
    if not root.is_dir():
        fail(f"Private root is not a directory: {root}")
    resolved = root.resolve()
    repo = ROOT.resolve()
    blocked = [repo / "release", repo / "build", repo / "chapters", repo / "signature"]
    for banned in blocked:
        if resolved == banned.resolve() or banned.resolve() in resolved.parents or resolved in banned.resolve().parents:
            if resolved != banned.resolve() and banned.resolve() in resolved.parents:
                fail(f"Refusing a private root inside {banned.name}.")
            if resolved == banned.resolve() or resolved in banned.resolve().parents:
                fail(f"Refusing to use {banned.name} as a private root.")
    if is_inside(resolved, repo):
        allowed = (repo / "journal" / "private").resolve()
        if not is_inside(resolved, allowed):
            fail("A private root inside the repository must stay under journal/private.")
    return resolved


def contained(path: Path, root: Path) -> Path:
    reject_symlink(path)
    if not is_inside(path, root):
        fail("Refusing a path outside the private root.")
    return path.resolve()


def read_yaml(path: Path) -> dict:
    contained_path = path
    try:
        loaded = yaml.safe_load(contained_path.read_text(encoding="utf-8-sig")) or {}
    except OSError as exc:
        fail(f"Could not read metadata: {exc.__class__.__name__}")
    if not isinstance(loaded, dict):
        fail("Session metadata must be a mapping.")
    return loaded


def load_sessions(private_root: Path) -> list[dict]:
    sessions_dir = private_root / "sessions"
    if not sessions_dir.exists():
        return []
    contained(sessions_dir, private_root)
    found = []
    for folder in sorted(path for path in sessions_dir.iterdir() if path.is_dir()):
        contained(folder, private_root)
        meta_path = folder / "session.yaml"
        if not meta_path.is_file():
            continue
        contained(meta_path, private_root)
        meta = read_yaml(meta_path)
        schema = str(meta.get("schema") or "")
        if schema and not schema.startswith("our-story-session/"):
            fail(f"Unsupported session schema in {folder.name}.")
        bodies = {}
        for key, name in BODY_FILES.items():
            file_path = folder / name
            if not file_path.exists():
                bodies[key] = ""
                continue
            contained(file_path, private_root)
            bodies[key] = file_path.read_text(encoding="utf-8-sig")
        mementos = []
        memento_path = folder / "mementos.yaml"
        if memento_path.exists():
            contained(memento_path, private_root)
            raw = read_yaml(memento_path)
            items = raw.get("items") if isinstance(raw.get("items"), list) else []
            for item in items:
                if not isinstance(item, dict):
                    continue
                filename = str(item.get("file") or "").strip()
                image = None
                missing = False
                if filename:
                    image_path = folder / filename
                    try:
                        reject_symlink(image_path)
                        if not is_inside(image_path, private_root):
                            fail("Refusing a memento path outside the private root.")
                        if image_path.is_file():
                            image = image_path.resolve()
                        else:
                            missing = True
                    except SystemExit:
                        raise
                    except OSError:
                        missing = True
                mementos.append(
                    {
                        "caption": str(item.get("caption") or "").strip(),
                        "image": image,
                        "missing": missing,
                    }
                )
        share = meta.get("share") if isinstance(meta.get("share"), dict) else {}
        found.append(
            {
                "folder_name": folder.name,
                "session_id": str(meta.get("session_id") or folder.name),
                "activity_id": str(meta.get("activity_id") or "").strip(),
                "title": str(meta.get("title") or "").strip(),
                "attempt": meta.get("attempt") or "",
                "date": str(meta.get("date") or ""),
                "time": str(meta.get("time") or ""),
                "place": str(meta.get("place") or ""),
                "version_played": str(meta.get("version_played") or ""),
                "status": str(meta.get("status") or ""),
                "content_revision": str(meta.get("content_revision") or ""),
                "share": {key: str(share.get(key) or "") for key in ("partner_a", "partner_b", "shared")},
                "prompts": meta.get("prompts") if isinstance(meta.get("prompts"), list) else [],
                "bodies": bodies,
                "mementos": mementos,
            }
        )
    found.sort(key=lambda item: (item["date"], str(item["attempt"]), item["activity_id"], item["session_id"]))
    return found


def include_body(mode: str, seat: str, share_flag: str) -> bool:
    if mode == "full-private":
        return True
    if mode == "shared":
        return share_flag == "shared"
    if seat == "shared":
        return share_flag == "shared"
    owner = "partner_a" if mode == "partner-a" else "partner_b"
    if seat == owner:
        return True
    return share_flag == "shared"


def visible_sections(session: dict, mode: str) -> list[tuple[str, str]]:
    sections = []
    for seat, label in (
        ("partner_a", "Partner A"),
        ("partner_b", "Partner B"),
        ("shared", "Shared"),
    ):
        if include_body(mode, seat, session["share"].get(seat, "")):
            sections.append((label, session["bodies"].get(seat) or ""))
    return sections


def render_html(sessions: list[dict], mode: str, stamp: str) -> str:
    parts = [
        "<!DOCTYPE html><html lang='en'><head><meta charset='utf-8'/>",
        "<title>Our Story</title>",
        "<style>",
        "@page { size: letter; margin: 0.85in 0.7in 0.75in 0.7in; }",
        "body { font-family: Georgia, serif; color: #241F1C; background: #FBF8F2; }",
        "h1 { font-family: 'Palatino Linotype', Palatino, serif; }",
        "article { break-after: page; }",
        "img { max-width: 4.5in; }",
        ".meta { font-family: Calibri, sans-serif; font-size: 11pt; }",
        "</style></head><body>",
        "<h1>Our Story</h1>",
        f"<p class='meta'>Local export · {html.escape(mode)} · {html.escape(stamp)}. FICTIONAL markers in a test are not a couple's memory.</p>",
        "<h2>Chronology</h2><ul>",
    ]
    for session in sessions:
        if not visible_sections(session, mode):
            continue
        label = " ".join(
            bit
            for bit in (
                session["date"],
                session["activity_id"],
                session["title"],
                f"attempt {session['attempt']}" if session["attempt"] != "" else "",
            )
            if bit
        )
        parts.append(f"<li>{html.escape(label)}</li>")
    parts.append("</ul>")
    for session in sessions:
        sections = visible_sections(session, mode)
        if not sections:
            continue
        parts.append("<article>")
        heading = " ".join(bit for bit in (session["activity_id"], session["title"]) if bit) or "Session"
        parts.append(f"<h2>{html.escape(heading)}</h2>")
        meta = " · ".join(
            bit
            for bit in (
                session["date"],
                session["time"],
                session["place"],
                session["version_played"],
                session["status"],
                f"attempt {session['attempt']}" if session["attempt"] != "" else "",
            )
            if bit
        )
        parts.append(f"<p class='meta'>{html.escape(meta)}</p>")
        for prompt in session["prompts"]:
            if not isinstance(prompt, dict):
                continue
            prompt_id = str(prompt.get("id") or "")
            wording = str(prompt.get("text_at_answer") or "")
            if prompt_id or wording:
                parts.append(
                    f"<p class='meta'>Prompt {html.escape(prompt_id)}: {html.escape(wording)}</p>"
                )
        for label, body in sections:
            parts.append(f"<h3>{html.escape(label)}</h3>")
            if body.strip():
                for paragraph in body.split("\n\n"):
                    text = paragraph.strip()
                    if text:
                        parts.append(f"<p>{html.escape(text)}</p>")
            else:
                parts.append("<p>This page was left blank.</p>")
        for item in session["mementos"]:
            parts.append("<h3>Keepsake</h3>")
            if item["caption"]:
                parts.append(f"<p>{html.escape(item['caption'])}</p>")
            if item["image"] is not None:
                parts.append(f"<img src='{item['image'].as_uri()}' alt='{html.escape(item['caption'] or 'Keepsake')}'/>")
            elif item["missing"]:
                parts.append("<p>The keepsake file was not in the session folder. The caption was kept.</p>")
        parts.append("</article>")
    parts.append("</body></html>")
    return "\n".join(parts)


def write_docx(path: Path, sessions: list[dict], mode: str, stamp: str) -> None:
    document = Document()
    document.add_heading("Our Story", 0)
    document.add_paragraph(f"Local export · {mode} · {stamp}")
    document.add_heading("Chronology", 1)
    for session in sessions:
        if not visible_sections(session, mode):
            continue
        bits = [session["date"], session["activity_id"], session["title"]]
        if session["attempt"] != "":
            bits.append(f"attempt {session['attempt']}")
        document.add_paragraph(" · ".join(bit for bit in bits if bit))
    for session in sessions:
        sections = visible_sections(session, mode)
        if not sections:
            continue
        heading = " ".join(bit for bit in (session["activity_id"], session["title"]) if bit) or "Session"
        document.add_heading(heading, 1)
        meta = " · ".join(
            bit
            for bit in (
                session["date"],
                session["place"],
                session["version_played"],
                f"attempt {session['attempt']}" if session["attempt"] != "" else "",
            )
            if bit
        )
        if meta:
            document.add_paragraph(meta)
        for prompt in session["prompts"]:
            if isinstance(prompt, dict) and (prompt.get("id") or prompt.get("text_at_answer")):
                document.add_paragraph(
                    f"Prompt {prompt.get('id') or ''}: {prompt.get('text_at_answer') or ''}".strip()
                )
        for label, body in sections:
            document.add_heading(label, 2)
            document.add_paragraph(body.strip() or "This page was left blank.")
        for item in session["mementos"]:
            document.add_heading("Keepsake", 2)
            if item["caption"]:
                document.add_paragraph(item["caption"])
            if item["image"] is not None:
                document.add_picture(str(item["image"]), width=Inches(4.5))
            elif item["missing"]:
                document.add_paragraph("The keepsake file was not in the session folder. The caption was kept.")
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Georgia"
            run.font.size = Pt(11)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def output_paths(private_root: Path, mode: str, stamp: str) -> tuple[Path, Path, Path]:
    folder = private_root / "exports"
    token = stamp.replace("-", "")
    stem = f"Our_Story_PRIVATE_MEMOIR_{token}_{mode}"
    return folder / f"{stem}.html", folder / f"{stem}.docx", folder / f"{stem}.pdf"


def export_memoir(private_root: Path, mode: str, stamp: str, authorize: set[str]) -> tuple[Path, Path]:
    root = assert_private_root(private_root)
    if mode not in MODES:
        fail(f"Unknown mode {mode}.")
    if mode == "full-private" and authorize != {"partner-a", "partner-b"}:
        fail("Full private export requires --authorize-partner-a and --authorize-partner-b.")
    sessions = load_sessions(root)
    html_path, docx_path, pdf_path = output_paths(root, mode, stamp)
    for path in (html_path, docx_path, pdf_path):
        contained(path.parent, root)
        if "PRIVATE_MEMOIR" not in path.name:
            fail("Refusing an export name that is not marked private.")
        if is_inside(path, ROOT / "release") or is_inside(path, ROOT / "build"):
            fail("Refusing to write an export into release or build.")
    page = render_html(sessions, mode, stamp)
    html_path.parent.mkdir(parents=True, exist_ok=True)
    html_path.write_text(page, encoding="utf-8")
    write_docx(docx_path, sessions, mode, stamp)
    build_all.print_pdf(html_path, pdf_path, error_path=html_path.with_suffix(".pdf-error.txt"), timeout=180)
    included = sum(1 for session in sessions if visible_sections(session, mode))
    print(f"mode: {mode}")
    print(f"sessions_read: {len(sessions)}")
    print(f"sessions_included: {included}")
    print(f"docx: {docx_path}")
    print(f"pdf: {pdf_path}")
    return docx_path, pdf_path


def parser() -> argparse.ArgumentParser:
    tool = argparse.ArgumentParser(description="Export a local Our Story memoir. Prints counts and paths only.")
    tool.add_argument("--private-root", type=Path, required=True)
    tool.add_argument("--mode", required=True, choices=MODES)
    tool.add_argument("--stamp", default=date.today().isoformat())
    tool.add_argument("--authorize-partner-a", action="store_true")
    tool.add_argument("--authorize-partner-b", action="store_true")
    return tool


def main(argv: list[str] | None = None) -> int:
    configure_stdio()
    args = parser().parse_args(argv)
    authorize = set()
    if args.authorize_partner_a:
        authorize.add("partner-a")
    if args.authorize_partner_b:
        authorize.add("partner-b")
    export_memoir(args.private_root, args.mode, args.stamp, authorize)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception as exc:
        raise SystemExit(f"Private memoir export failed: {exc.__class__.__name__}") from exc
