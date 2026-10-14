#!/usr/bin/env python3
"""Build the Our Story companion journal as DOCX and PDF.

Usage, from the repository root:

    python scripts/build_journal.py

Reads journal/COMPANION_STRUCTURE.md, journal/PROMPT_BANK.md, and
journal/templates/*.md. Writes build/journal/our-story.html,
build/journal/Our_Story_Companion.docx, and
build/journal/Our_Story_Companion.pdf. Print CSS is styles/journal.css,
inlined into the HTML. The PDF uses the same Chrome or Edge headless
print as scripts/build_all.py.

Does not read journal/private/. Does not fill in personal answers.
Copy the source sheets as written. Do not invent personal answers or
fictional memories. If Chrome and Edge cannot print, the HTML and DOCX
are kept and the error is written to build/journal/pdf-print-error.txt.
"""

from __future__ import annotations

import json
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_all

ROOT = build_all.ROOT
OUT_DIR = ROOT / "build" / "journal"
HTML_PATH = OUT_DIR / "our-story.html"
DOCX_PATH = OUT_DIR / "Our_Story_Companion.docx"
PDF_PATH = OUT_DIR / "Our_Story_Companion.pdf"
PDF_ERROR_PATH = OUT_DIR / "pdf-print-error.txt"
CSS_PATH = ROOT / "styles" / "journal.css"
CATALOG_PATH = ROOT / "source" / "experience_catalog.json"
TABLE_TWIPS = 10008
RULE_TOKEN = "{{RULE}}"
STATEMENT = (
    "A personal companion to The Closer Project. The activities live in the main volume. "
    "This book is the record you two choose to keep."
)
PACK_SENTENCE = (
    "The repeatable pack is the session log, the partner page, the shared page, "
    "and the then-versus-now sheet, printed without a specific ID filled in."
)
HAND_SENTENCE = "Unlimited copies. The couple writes the ID, the title, and the attempt number by hand."
CROSS_SENTENCE = "All 108 IDs in catalog order, with part title, activity title, and the special sheet if one exists."
BANK_PROMPTS = ("**Partner A:", "**Partner B:", "**Shared:", "**Revisit:")
TEMPLATE_ORDER = (
    "session-log.md",
    "partner-page.md",
    "shared-page.md",
    "revisit-then-now.md",
    "memento-caption.md",
    "chronology-index.md",
    "milestone-page.md",
    "letter-page.md",
    "postcard.md",
    "yes-no-maybe.md",
    "hundred-points.md",
    "museum-label.md",
    "escape-clue-sheet.md",
    "civilization-charter.md",
)
SPECIAL_ROW = re.compile(r"^\|\s*`([^`]+)`\s*\|\s*([^|]+)\|")


def allow(path) -> None:
    build_all.assert_readable(path)
    rel = path.resolve().relative_to(ROOT.resolve()).as_posix()
    permitted = (
        rel in {"journal/COMPANION_STRUCTURE.md", "journal/PROMPT_BANK.md", "source/experience_catalog.json"}
        or rel.startswith("journal/templates/")
        or rel == "styles/journal.css"
    )
    if not permitted:
        raise SystemExit(f"Journal build refuses to read {rel}")


def read_text(path) -> str:
    allow(path)
    return path.read_text(encoding="utf-8-sig")


def require_phrase(text: str, phrase: str) -> None:
    if phrase not in text:
        raise SystemExit(f"Missing companion wording from the source files: {phrase}")


def preprocess_rules(text: str) -> str:
    out: list[str] = []
    in_fence = False
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("```"):
            in_fence = not in_fence
            out.append(line)
            continue
        if in_fence or not line.startswith(">"):
            out.append(line)
            continue
        content = line[1:].lstrip()
        out.append("")
        if content.strip() and set(content.strip()) <= {"_"}:
            out.append(RULE_TOKEN)
        elif content.strip():
            out.append(content)
        out.append("")
    return "\n".join(out)


def is_bank_prompt(text: str) -> bool:
    stripped = text.strip()
    return any(stripped.startswith(label) for label in BANK_PROMPTS)


def blocks_from_markdown(text: str, *, expand_prompts: bool = False, rules_from_quotes: bool = False, skip_fences: bool = False) -> list:
    if rules_from_quotes:
        text = preprocess_rules(text)
    parsed = build_all.parse_blocks(text, "journal")
    blocks = []
    for block in parsed:
        kind = block[0]
        if kind == "heading":
            blocks.append(("heading", block[1], block[2]))
        elif kind == "para":
            if block[1].strip() == RULE_TOKEN:
                blocks.append(("rule", 1))
                continue
            prompt = expand_prompts and is_bank_prompt(block[1])
            blocks.append(("para", block[1], prompt))
            if prompt:
                blocks.append(("rule", 8))
        elif kind == "list":
            blocks.append(("list", block[1], block[2]))
        elif kind == "table":
            blocks.append(("table", block[1], block[2]))
        elif kind == "fence":
            if not skip_fences:
                blocks.append(("pre", block[2]))
        elif kind == "hr":
            blocks.append(("hr",))
    return blocks


def special_sheets(structure: str) -> dict[str, str]:
    found: dict[str, list[str]] = {}
    for line in structure.splitlines():
        match = SPECIAL_ROW.match(line)
        if not match:
            continue
        sheet = match.group(1).strip()
        for raw in re.findall(r"\b(\d{3})\b", match.group(2)):
            bucket = found.setdefault(raw, [])
            if sheet not in bucket:
                bucket.append(sheet)
    return {key: ", ".join(names) for key, names in found.items()}


def catalog_rows() -> list[dict]:
    allow(CATALOG_PATH)
    data = json.loads(CATALOG_PATH.read_text(encoding="utf-8"))
    rows = []
    for item in data.get("experiences") or []:
        number = int(item["id"])
        rows.append(
            {
                "id": f"{number:03d}",
                "part": str(item.get("part_title") or ""),
                "title": str(item.get("title") or ""),
            }
        )
    rows.sort(key=lambda row: int(row["id"]))
    return rows


def template_paths() -> list:
    folder = ROOT / "journal" / "templates"
    if not folder.is_dir():
        raise SystemExit(f"Missing journal templates: {folder}")
    found = {}
    for path in folder.glob("*.md"):
        if path.is_file():
            allow(path)
            found[path.name] = path
    ordered = []
    for name in TEMPLATE_ORDER:
        if name in found:
            ordered.append(found.pop(name))
    for name in sorted(found):
        ordered.append(found[name])
    if not ordered:
        raise SystemExit("No journal templates were found. No companion files were written.")
    return ordered


def title_blocks(structure: str) -> list:
    require_phrase(structure, STATEMENT)
    require_phrase(structure, "A companion journal")
    require_phrase(structure, "names or initials")
    return [
        ("kicker", "OUR STORY"),
        ("heading", 1, "Our Story"),
        ("statement", "A companion journal."),
        ("statement", STATEMENT),
        ("label", "Name or initials"),
        ("rule", 1),
        ("label", "Name or initials"),
        ("rule", 1),
    ]


def body_blocks(structure: str, bank: str) -> list:
    require_phrase(structure, "Eight handwriting rules")
    require_phrase(structure, "Date we began this companion")
    require_phrase(structure, PACK_SENTENCE)
    require_phrase(structure, HAND_SENTENCE)
    blocks: list = []
    blocks.append(("heading", 1, "How to use these pages"))
    blocks.append(("para", "Date what you write. Skip what you do not want stored. Private pages are optional to share.", False))
    blocks.append(("para", "A written yes is not a promise to act, and a game result is not consent to touch.", False))
    blocks.append(("para", "No real memory is supplied here. The lines stay blank until you fill them.", False))
    blocks.append(("heading", 1, "Dedication"))
    blocks.append(("rule", 8))
    blocks.append(("heading", 1, "Date we began this companion"))
    blocks.append(("label", "Place"))
    blocks.append(("rule", 1))
    blocks.append(("label", "Initials, if wanted"))
    blocks.append(("rule", 1))
    blocks.append(("label", "First sentence, if you want one"))
    blocks.append(("rule", 8))
    rows = catalog_rows()
    sheets = special_sheets(structure)
    blocks.append(("heading", 1, "Cross-reference"))
    if len(rows) == 108 and CROSS_SENTENCE in structure:
        blocks.append(("para", CROSS_SENTENCE, False))
    else:
        blocks.append(("para", f"The catalog lists {len(rows)} experiences.", False))
    table_rows = [
        [row["id"], row["part"], row["title"], sheets.get(row["id"], "")]
        for row in rows
    ]
    blocks.append(("table", ["ID", "Part", "Title", "Special sheet"], table_rows))
    blocks.append(("heading", 1, "Sheets"))
    blocks.append(("para", PACK_SENTENCE, False))
    blocks.append(("para", HAND_SENTENCE, False))
    for path in template_paths():
        blocks.extend(blocks_from_markdown(read_text(path), rules_from_quotes=True, skip_fences=True))
    blocks.extend(blocks_from_markdown(bank[bank.find("\n## Part "):] if "\n## Part " in bank else bank, expand_prompts=True, skip_fences=True))
    blocks.append(("heading", 1, "A note"))
    blocks.append(("para", "These pages are blank on purpose. Write only what you both choose to keep.", False))
    return blocks


def html_blocks(blocks: list) -> str:
    parts: list[str] = []
    index = 0
    while index < len(blocks):
        block = blocks[index]
        kind = block[0]
        if kind == "kicker":
            parts.append(f'<p class="kicker">{build_all.esc(block[1])}</p>')
        elif kind == "heading":
            level = min(max(int(block[1]), 1), 4)
            parts.append(f"<h{level}>{build_all.esc(block[2])}</h{level}>")
        elif kind == "statement":
            parts.append(f'<p class="statement">{build_all.esc(block[1])}</p>')
        elif kind == "label":
            parts.append(f"<p>{build_all.esc(block[1])}</p>")
        elif kind == "para":
            text = block[1]
            css = "fact-line" if "____" in text else ""
            if len(block) > 2 and block[2]:
                parts.append('<div class="prompt-block">')
                parts.append(f"<p>{build_all.inline_xhtml(text, None)}</p>")
                if index + 1 < len(blocks) and blocks[index + 1][0] == "rule":
                    index += 1
                    parts.append(build_all.writing_lines(blocks[index][1]))
                parts.append("</div>")
            else:
                class_attr = f' class="{css}"' if css else ""
                parts.append(f"<p{class_attr}>{build_all.inline_xhtml(text, None)}</p>")
        elif kind == "rule":
            parts.append(build_all.writing_lines(block[1]))
        elif kind == "list":
            parts.append(build_all.list_xhtml(block[1], block[2]))
        elif kind == "table":
            parts.append(build_all.table_xhtml(block[1], block[2]))
        elif kind == "pre":
            parts.append(f"<pre><code>{build_all.esc(block[1])}</code></pre>")
        elif kind == "hr":
            parts.append("<hr/>")
        index += 1
    return "\n".join(parts)


def html_document(css: str, title: list, body: list) -> str:
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<meta name="generator" content="scripts/build_journal.py"/>
<title>Our Story — a companion journal</title>
<style>
{css}
</style>
</head>
<body>
<header class="title-page">
{html_blocks(title)}
</header>
{html_blocks(body)}
</body>
</html>
"""


def configure_document(document: Document) -> None:
    build_all.configure_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.70)
    section.top_margin = Inches(0.70)
    section.bottom_margin = Inches(0.75)
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.32)
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("OUR STORY")
    build_all.set_run_font(run, "Palatino Linotype", 8, build_all.COPPER)
    section.first_page_header.paragraphs[0].text = ""
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    build_all.add_page_field(footer)
    section.first_page_footer.paragraphs[0].text = ""
    background = build_all.OxmlElement("w:background")
    background.set(build_all.qn("w:color"), "F6F1E7")
    document.element.insert(0, background)
    document.settings.element.append(build_all.OxmlElement("w:displayBackgroundShape"))
    props = document.core_properties
    props.title = "Our Story — a companion journal"
    props.subject = "Companion to The Closer Project"
    props.author = "The Closer Project"
    props.category = "Companion"
    props.comments = "Generated from blank journal templates by scripts/build_journal.py. No personal answers."
    props.last_modified_by = "The Closer Project"


def add_rules(document: Document, count: int) -> None:
    for _ in range(count):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(26)
        run = paragraph.add_run(" ")
        build_all.set_run_font(run, "Georgia", 11, build_all.INK)
        build_all.add_bottom_border(paragraph, "D9CFC0")


def add_table(document: Document, header: list[str], rows: list[list[str]], bookmarks) -> None:
    table = document.add_table(rows=1 + len(rows), cols=max(1, len(header)))
    build_all.set_table_width(table, TABLE_TWIPS)
    build_all.set_table_borders(table, "D9CFC0", sz="4", inside=True)
    for index, text in enumerate(header):
        cell = table.cell(0, index)
        build_all.shade_cell(cell, "F6F1E7")
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(2)
        build_all.add_inlines(paragraph, text, bookmarks, "Palatino Linotype", 11, build_all.INK)
    for r_index, row in enumerate(rows):
        padded = list(row) + [""] * max(0, len(header) - len(row))
        for c_index, text in enumerate(padded[: len(header)]):
            cell = table.cell(r_index + 1, c_index)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            build_all.add_inlines(paragraph, str(text), bookmarks, "Georgia", 10, build_all.INK)
    document.add_paragraph().paragraph_format.space_after = Pt(6)


def render_docx(document: Document, blocks: list, bookmarks, *, break_before_h1: bool, center_title: bool) -> None:
    first_h1 = True
    index = 0
    while index < len(blocks):
        block = blocks[index]
        kind = block[0]
        if kind == "heading" and int(block[1]) == 1:
            if break_before_h1 or not first_h1:
                document.add_page_break()
            first_h1 = False
        if kind == "kicker":
            build_all.add_text_p(
                document,
                block[1],
                bookmarks,
                font="Calibri",
                size=9,
                color=build_all.COPPER,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                before=48,
                after=6,
            )
        elif kind == "heading":
            level = int(block[1])
            style = {1: "Heading 1", 2: "Heading 2"}.get(level, "Heading 3")
            size = {1: 28, 2: 16}.get(level, 13)
            color = build_all.BURGUNDY if level == 2 else build_all.INK
            align = WD_ALIGN_PARAGRAPH.CENTER if center_title and level == 1 and not break_before_h1 else None
            build_all.add_text_p(
                document,
                block[2],
                bookmarks,
                style=style,
                font="Palatino Linotype",
                size=size,
                color=color,
                align=align,
                before=12 if level != 1 else 0,
                after=6,
                keep=True,
            )
        elif kind in {"statement", "label", "para"}:
            text = block[1]
            italic = kind == "statement"
            align = WD_ALIGN_PARAGRAPH.CENTER if center_title and kind == "statement" else None
            prompt = kind == "para" and len(block) > 2 and block[2]
            fact = kind != "statement" and "____" in text
            if fact:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(26)
                if align is not None:
                    paragraph.alignment = align
                build_all.add_inlines(paragraph, text, bookmarks, "Georgia", 11, build_all.INK, italic=italic)
            else:
                build_all.add_text_p(
                    document,
                    text,
                    bookmarks,
                    italic=italic,
                    align=align,
                    after=6,
                    keep=bool(prompt),
                )
            if prompt and index + 1 < len(blocks) and blocks[index + 1][0] == "rule":
                index += 1
                add_rules(document, blocks[index][1])
        elif kind == "rule":
            add_rules(document, block[1])
        elif kind == "list":
            style = "List Number" if block[1] else "List Bullet"
            for item in block[2]:
                build_all.add_text_p(document, item, bookmarks, style=style, after=3)
        elif kind == "table":
            add_table(document, block[1], block[2], bookmarks)
        elif kind == "pre":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block[1])
            build_all.set_run_font(run, "Consolas", 9, build_all.INK)
        elif kind == "hr":
            paragraph = document.add_paragraph()
            build_all.add_bottom_border(paragraph, "A85A32")
        index += 1


def write_docx(title: list, body: list) -> None:
    document = Document()
    configure_document(document)
    bookmarks = build_all.Bookmarks()
    render_docx(document, title, bookmarks, break_before_h1=False, center_title=True)
    render_docx(document, body, bookmarks, break_before_h1=True, center_title=False)
    tmp = DOCX_PATH.with_name(DOCX_PATH.name + ".tmp")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(tmp)
    build_all.normalize_office_zip(tmp, DOCX_PATH)


def guard_private() -> None:
    try:
        build_all.assert_readable(ROOT / "journal" / "private" / "unused.md")
    except SystemExit as exc:
        if "Refusing" not in str(exc):
            raise
    else:
        raise SystemExit("Private-path guard did not refuse journal/private.")
    if build_all.is_private_path(OUT_DIR):
        raise SystemExit("Refusing to write the companion into journal/private.")


def main() -> int:
    build_all.configure_stdio()
    guard_private()
    if not CSS_PATH.is_file():
        raise SystemExit(f"Missing stylesheet: {CSS_PATH}")
    structure_path = ROOT / "journal" / "COMPANION_STRUCTURE.md"
    bank_path = ROOT / "journal" / "PROMPT_BANK.md"
    structure = read_text(structure_path)
    bank = read_text(bank_path)
    title = title_blocks(structure)
    body = body_blocks(structure, bank)
    css = read_text(CSS_PATH)
    page = html_document(css, title, body)
    if "DO_NOT_COMMIT" in page:
        raise SystemExit("Companion HTML picked up a private marker. No files were written.")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_all.write_bytes(HTML_PATH, page.encode("utf-8"))
    write_docx(title, body)
    try:
        build_all.print_pdf(
            HTML_PATH,
            PDF_PATH,
            error_path=PDF_ERROR_PATH,
            kept="The companion HTML and DOCX were kept.",
        )
    except SystemExit:
        for path in (HTML_PATH, DOCX_PATH):
            if path.is_file():
                print(f"KEPT  {path.relative_to(ROOT).as_posix()}  {path.stat().st_size} bytes")
        raise
    if PDF_ERROR_PATH.exists():
        PDF_ERROR_PATH.unlink()
    for path in (HTML_PATH, DOCX_PATH, PDF_PATH):
        print(f"WROTE  {path.relative_to(ROOT).as_posix()}  {path.stat().st_size} bytes")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception as exc:
        raise SystemExit(f"Journal build failed: {exc}") from exc
